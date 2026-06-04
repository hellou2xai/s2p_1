"""Augment the two big training guides so every worked example carries a
folder-tree code block and a 'What Claude did, behind the scenes' walkthrough,
as required by CLAUDE.md.

Strategy: open each .docx, walk paragraphs, find the marker line for each
known use case or vignette, and insert new paragraphs immediately after it
using XML addnext. Existing content is preserved; we only inject.
"""

from __future__ import annotations

import shutil
import time
from pathlib import Path
from typing import Iterable, List

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.paragraph import Paragraph


ROOT = Path(
    r"C:\Users\sambi\OneDrive - U2xAI\Claude Code Projects"
    r"\S2p Training\Course List\Claude Code Foundation for S2P"
)
HANDOUTS = ROOT / "Handouts"


# ---------------------------------------------------------------------------
# XML helpers (insert paragraphs after a target paragraph in document order)
# ---------------------------------------------------------------------------

def _new_p_after(prev_p: Paragraph) -> Paragraph:
    new_p = OxmlElement("w:p")
    prev_p._element.addnext(new_p)
    return Paragraph(new_p, prev_p._parent)


def insert_para_after(prev: Paragraph, text: str, *, bold=False) -> Paragraph:
    p = _new_p_after(prev)
    run = p.add_run(text)
    if bold:
        run.bold = True
    return p


def insert_code_after(prev: Paragraph, text: str) -> Paragraph:
    p = _new_p_after(prev)
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(10)
    return p


def insert_numbered_list_after(prev: Paragraph, items: List[str]) -> Paragraph:
    last = prev
    for i, item in enumerate(items, start=1):
        p = _new_p_after(last)
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(3)
        p.add_run(f"{i}. {item}")
        last = p
    return last


def find_paragraph_by_text(doc, fragment: str) -> Paragraph | None:
    fragment = fragment.strip().lower()
    for p in doc.paragraphs:
        if fragment in p.text.strip().lower():
            return p
    return None


def find_paragraph_after(doc, anchor_fragment: str, after_fragment: str) -> Paragraph | None:
    """Return the first paragraph containing after_fragment that comes
    after the paragraph containing anchor_fragment."""
    anchor = find_paragraph_by_text(doc, anchor_fragment)
    if anchor is None:
        return None
    found_anchor = False
    for p in doc.paragraphs:
        if p._element is anchor._element:
            found_anchor = True
            continue
        if found_anchor and after_fragment.lower() in p.text.lower():
            return p
    return None


def open_with_retry(path: Path):
    for _ in range(6):
        try:
            return Document(str(path))
        except PermissionError:
            time.sleep(0.5)
    raise PermissionError(f"could not open {path}; close Word or pause OneDrive sync")


# ---------------------------------------------------------------------------
# Augmentation specs
# ---------------------------------------------------------------------------

# Each entry is one worked example in a doc. It says:
#   anchor: a unique fragment from the heading or first paragraph of the example
#   insert_after: a unique fragment of the paragraph AFTER which we insert
#                  (typically the "What Cowork produces" / result paragraph)
#   folder: an ASCII folder tree (multi-line string)
#   behind_scenes: the numbered walkthrough steps

CODE_DESKTOP_AUGMENTS = [
    # ------------------ Day in the Life vignettes ------------------
    {
        "anchor": "7:00 AM: Reviewing Overnight Routine Outputs",
        "insert_after": "Total elapsed time",
        "folder": (
            "Procurement_Analytics/\n"
            "├── routines/\n"
            "│   └── nightly_oracle_pull.py\n"
            "├── data/\n"
            "│   └── oracle_exports/\n"
            "│       └── spend_2026-04-25.csv\n"
            "├── outputs/\n"
            "│   └── overnight_briefings/\n"
            "│       └── 2026-04-25_brief.md\n"
            "└── CLAUDE.md"
        ),
        "behind_scenes": [
            "At 02:00 the scheduled Routine ran nightly_oracle_pull.py against the live Oracle export endpoint.",
            "Pulled the day's PO and invoice data into data/oracle_exports/spend_2026-04-25.csv.",
            "Compared the new spend rows against the previous seven days' rolling baseline.",
            "Flagged any commodity code where today's spend was more than 20% above the seven-day average.",
            "Wrote the flags as a one-page Markdown briefing to outputs/overnight_briefings/.",
            "Sent a Slack DM with the briefing link so it was waiting for the procurement manager at 07:00.",
        ],
    },
    {
        "anchor": "9:30 AM: Running Parallel Supplier Scorecards",
        "insert_after": "session sidebar",
        "folder": (
            "Procurement_Analytics/\n"
            "├── data/\n"
            "│   └── kpi_exports/q1_kpis.csv\n"
            "├── templates/\n"
            "│   └── scorecard.xlsx\n"
            "└── outputs/\n"
            "    └── scorecards/\n"
            "        ├── electronics/\n"
            "        ├── packaging/\n"
            "        └── logistics/"
        ),
        "behind_scenes": [
            "Three parallel sessions opened in the multi-session sidebar, each on its own Git worktree branch.",
            "Session A loaded electronics suppliers from kpi_exports/q1_kpis.csv. Session B loaded packaging. Session C loaded logistics.",
            "Each session ran the same scorecard build script against its own subset of suppliers.",
            "Because each session has an isolated Git branch, the three runs cannot overwrite each other's outputs.",
            "Each session saved to its own outputs/scorecards/<category>/ folder.",
            "Status indicators in the sidebar showed 'running', 'running', 'completed' so the manager could see progress at a glance.",
        ],
    },
    {
        "anchor": "11:00 AM: Processing a Contract Batch",
        "insert_after": "deviation table",
        "folder": (
            "Procurement_Analytics/\n"
            "├── contracts/\n"
            "│   └── inbound_april/   (12 supplier MSAs as PDF)\n"
            "├── standards/\n"
            "│   └── our_standard_msa.docx\n"
            "└── outputs/\n"
            "    └── contract_reviews/\n"
            "        ├── deviations.xlsx\n"
            "        └── per_contract_summaries/"
        ),
        "behind_scenes": [
            "Listed every PDF in contracts/inbound_april/. Counted 12.",
            "Read standards/our_standard_msa.docx as the comparison baseline.",
            "For each PDF, OCR'd if the document was scanned, then parsed clauses by numbering pattern.",
            "Compared each clause against the matching baseline clause and flagged any deviation.",
            "Wrote one row per deviation to outputs/contract_reviews/deviations.xlsx with clause, contract, deviation, severity.",
            "Wrote a per-contract executive summary to outputs/contract_reviews/per_contract_summaries/.",
            "Master/standards files were never modified.",
        ],
    },
    {
        "anchor": "2:00 PM: Dispatching a Task from His Phone",
        "insert_after": "API",
        "folder": (
            "On the phone:\n"
            "  Shortcut: 'New supplier risk brief'\n"
            "  Triggers: POST https://routines.claude.com/<id>\n"
            "\n"
            "On the laptop, when Claude Code Desktop wakes:\n"
            "  Procurement_Analytics/\n"
            "  └── outputs/\n"
            "      └── ad_hoc_briefs/\n"
            "          └── techsource_2026-04-25.md"
        ),
        "behind_scenes": [
            "The phone shortcut sent a POST request to the API Routine's unique HTTP endpoint with a small JSON payload (the supplier name).",
            "Claude Code Desktop received the trigger, opened a fresh session, and read the supplier name from the payload.",
            "Pulled the most recent spend, contract, and risk data for that supplier from Oracle export CSVs.",
            "Pulled the latest news on the supplier via the connected web-search tool.",
            "Drafted a one-page risk brief and saved it to outputs/ad_hoc_briefs/.",
            "Returned a download link in the API response, which the phone shortcut surfaced as a notification.",
        ],
    },
    {
        "anchor": "4:30 PM: Setting Up Tomorrow's Routine",
        "insert_after": "Routine",
        "folder": (
            "Procurement_Analytics/\n"
            "├── routines/\n"
            "│   └── savings_tracker_weekly.py\n"
            "├── data/\n"
            "│   └── savings_register.xlsx\n"
            "└── outputs/\n"
            "    └── savings_reports/"
        ),
        "behind_scenes": [
            "Defined a new scheduled Routine to run every Friday at 06:00.",
            "The Routine reads data/savings_register.xlsx and the most recent Oracle spend export.",
            "It compares actuals against the negotiated baselines for each savings initiative.",
            "It computes realised vs planned savings and flags any initiative that has slipped by more than 5%.",
            "It saves a weekly savings report to outputs/savings_reports/<date>.md.",
            "It emails the report to the procurement leadership distribution list.",
        ],
    },
    # ------------------ Use Cases ------------------
    {
        "anchor": "Use Case 1: Oracle Spend Variance Analysis from CSV Export",
        "insert_after": "variance",
        "folder": (
            "Procurement_Analytics/\n"
            "├── data/\n"
            "│   └── oracle_exports/\n"
            "│       ├── purchasing_q3.csv\n"
            "│       └── q3_approved_budget.xlsx\n"
            "├── reference/\n"
            "│   └── commodity_crosswalk.csv\n"
            "└── outputs/\n"
            "    └── variance_q3.xlsx"
        ),
        "behind_scenes": [
            "Read data/oracle_exports/purchasing_q3.csv (the actuals) and data/oracle_exports/q3_approved_budget.xlsx (the budget).",
            "Normalised supplier names against the Oracle supplier master so duplicates merged.",
            "Joined the two datasets on commodity code, using reference/commodity_crosswalk.csv to map departments to commodities.",
            "Calculated variance per commodity (actual minus budget, both absolute and percentage).",
            "Flagged every commodity where variance is more than 10% above budget.",
            "Wrote outputs/variance_q3.xlsx with one row per commodity and a summary tab listing the top ten variances.",
            "Did not modify anything in data/.",
        ],
    },
    {
        "anchor": "Use Case 2: Multi-Supplier Scorecard Automation",
        "insert_after": "scorecards",
        "folder": (
            "Procurement_Analytics/\n"
            "├── data/\n"
            "│   └── kpi_export.csv\n"
            "├── templates/\n"
            "│   └── scorecard_master.xlsx\n"
            "└── outputs/\n"
            "    └── scorecards/\n"
            "        └── (50 .xlsx files, one per supplier)"
        ),
        "behind_scenes": [
            "Loaded data/kpi_export.csv. Identified 50 distinct suppliers.",
            "Opened templates/scorecard_master.xlsx. Worked out where each KPI goes by reading the cell labels.",
            "For each supplier, created a copy of the template and filled in their four KPIs (delivery, quality, price compliance, invoice accuracy).",
            "Computed RAG status per KPI using the bands defined in CLAUDE.md.",
            "Wrote a one-line commentary on each supplier based on their KPI mix.",
            "Saved one .xlsx per supplier to outputs/scorecards/ named after the supplier.",
            "templates/ and data/ were never modified.",
        ],
    },
    {
        "anchor": "Use Case 3: Contract Clause Batch Extraction",
        "insert_after": "clause",
        "folder": (
            "Procurement_Analytics/\n"
            "├── contracts/\n"
            "│   └── batch_april/   (10 supplier contracts)\n"
            "└── outputs/\n"
            "    └── clause_extraction.xlsx"
        ),
        "behind_scenes": [
            "Listed every file in contracts/batch_april/. Counted 10.",
            "For each file, read the document text. OCR'd any scanned PDFs.",
            "Searched each contract for the named clauses: limitation of liability, termination, indemnity, IP ownership, data protection, governing law, auto-renewal.",
            "Recorded the clause text and the page number for each match.",
            "Wrote one row per (contract, clause) pair to outputs/clause_extraction.xlsx.",
            "Where a clause was missing, recorded 'not present' rather than guessing.",
            "contracts/ files were not modified.",
        ],
    },
    {
        "anchor": "Use Case 4: Savings Tracking Dashboard",
        "insert_after": "dashboard",
        "folder": (
            "Procurement_Analytics/\n"
            "├── data/\n"
            "│   ├── savings_register.xlsx\n"
            "│   └── oracle_exports/\n"
            "│       └── ytd_spend.csv\n"
            "└── outputs/\n"
            "    ├── savings_dashboard.html\n"
            "    └── savings_summary.md"
        ),
        "behind_scenes": [
            "Read data/savings_register.xlsx (the planned-savings register, one row per initiative).",
            "Read data/oracle_exports/ytd_spend.csv (year-to-date actual spend).",
            "For each initiative, joined planned baseline against actual spend on the matched commodity code.",
            "Calculated realised savings and percentage of plan achieved per initiative.",
            "Built an interactive HTML dashboard with one tile per initiative, RAG-coloured by % achieved.",
            "Wrote a one-page Markdown summary listing the three best and three worst initiatives with figures.",
            "Saved both files to outputs/. Did not modify data/.",
        ],
    },
]


COWORK_AUGMENTS = [
    # ------------------ Day in the Life vignettes ------------------
    {
        "anchor": "7:45 AM: Preparing for a Supplier QBR",
        "insert_after": "scorecard",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   ├── about-me.md\n"
            "│   ├── my-company.md\n"
            "│   └── procurement-standards.md\n"
            "├── TEMPLATES/\n"
            "│   ├── scorecard.docx\n"
            "│   └── talking_points.docx\n"
            "└── OUTPUTS/\n"
            "    └── apex-materials-q3-qbr/\n"
            "        ├── scorecard.docx\n"
            "        └── talking-points.docx"
        ),
        "behind_scenes": [
            "Read every file in ABOUT ME/ so it knew Sarah's role, her company's targets, and the procurement scoring rules.",
            "Asked three clarifying questions because the prompt was non-routine: target audience, three things to recognise, three things to put on notice.",
            "Loaded the four KPIs from the prompt and applied the RAG bands from procurement-standards.md.",
            "Filled TEMPLATES/scorecard.docx with the KPI table, RAG-coloured.",
            "Drafted talking points opening with the 99.1% quality (positive recognition), pivoting to the three October delivery slips, closing with the action ask.",
            "Saved both files into a new dated subfolder under OUTPUTS/.",
            "Did not write outside OUTPUTS/.",
        ],
    },
    {
        "anchor": "9:30 AM: Flagging a Contract Clause Deviation",
        "insert_after": "deviation table",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   └── procurement-standards.md\n"
            "└── OUTPUTS/\n"
            "    └── contract-reviews/\n"
            "        ├── chem-supplier-msa-v2.pdf            (input)\n"
            "        └── chem-supplier-msa-v2-review.md     (output)"
        ),
        "behind_scenes": [
            "Opened OUTPUTS/contract-reviews/chem-supplier-msa-v2.pdf and read it end to end.",
            "Loaded ABOUT ME/procurement-standards.md to know the company's standard contract terms.",
            "Walked the contract clause by clause and compared each against the matching standard.",
            "Flagged the liability cap as $500K below standard, the audit rights window as 90 days versus 180, and the indemnity as one-way against the supplier instead of mutual.",
            "Built a deviation table sorted by risk (highest first) with clause, deviation, risk level, our standard, recommended fix.",
            "Wrote a three-bullet executive summary at the top of the same Markdown file.",
            "Saved as chem-supplier-msa-v2-review.md alongside the source PDF.",
        ],
    },
    {
        "anchor": "11:00 AM: Building a Savings Case for the CFO",
        "insert_after": "memo",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   └── my-company.md\n"
            "└── OUTPUTS/\n"
            "    ├── spend-reports/\n"
            "    │   └── logistics-q1-q3-spend.csv\n"
            "    └── savings-cases/\n"
            "        └── logistics-consolidation.md"
        ),
        "behind_scenes": [
            "Read OUTPUTS/spend-reports/logistics-q1-q3-spend.csv to verify the $11.2M baseline against the 14 incumbents.",
            "Asked four clarifying questions about transition costs, cut-over risk, the CFO's known soft-savings scepticism, and existing volume commitments.",
            "Computed hard savings as baseline minus negotiated rate ($11.2M minus $9.1M = $2.1M annually).",
            "Computed soft savings of $380K from reduced supplier-management overhead. Labelled clearly as soft.",
            "Estimated transition costs and offset them against year-one savings.",
            "Built three risk-adjusted scenarios (best, base, downside) with their savings figures.",
            "Wrote a stand-alone objection-handling section addressing the 'is this real savings' question with concrete proof points.",
            "Saved as a two-page executive memo to OUTPUTS/savings-cases/logistics-consolidation.md.",
        ],
    },
    {
        "anchor": "2:00 PM: Running a Quick Supplier Risk Check",
        "insert_after": "brief",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "└── OUTPUTS/\n"
            "    └── supplier-scorecards/\n"
            "        └── techsource-risk-brief.md"
        ),
        "behind_scenes": [
            "Read the prompt: single-source supplier, gallium shortage, Q4 demand 45,000 units, last delivery 8,000 units, audience VP Operations, deadline 45 minutes.",
            "Computed the supply gap (Q4 demand minus last confirmed delivery = 37,000 unit shortfall).",
            "Wrote a three-sentence risk summary leading with the supply gap.",
            "Researched alternative chip suppliers using the connected web-search tool, focusing on those with declared MCU stock.",
            "Listed three alternatives with rough lead time and qualification effort estimates.",
            "Ranked recommended actions by speed of impact: emergency stock release first, qualified second source second, redesign for an alternate chip third.",
            "Saved the brief to OUTPUTS/supplier-scorecards/techsource-risk-brief.md within the 60-second response.",
        ],
    },
    {
        "anchor": "4:30 PM: Saving a Template Before Closing",
        "insert_after": "TEMPLATES",
        "folder": (
            "Claude Cowork/\n"
            "├── TEMPLATES/\n"
            "│   └── pre_negotiation_brief_template.md  (newly saved)\n"
            "└── OUTPUTS/\n"
            "    └── apex-materials-q3-qbr/\n"
            "        └── talking-points.docx           (the source)"
        ),
        "behind_scenes": [
            "Read OUTPUTS/apex-materials-q3-qbr/talking-points.docx (the doc Sarah was happy with at 09:00).",
            "Stripped the supplier-specific facts (Apex name, October dates, the 94.2% delivery figure) and replaced them with placeholder tokens like {SupplierName} and {DeliveryActualPct}.",
            "Kept the structural skeleton: opening recognition, three middle points, action ask, close.",
            "Saved the cleaned skeleton to TEMPLATES/pre_negotiation_brief_template.md.",
            "Did not change OUTPUTS/. The original talking points stay as the worked example for next time.",
        ],
    },
    # ------------------ Use Cases ------------------
    {
        "anchor": "Use Case 1: RFP Package Development",
        "insert_after": "supplier-response-template.md",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   ├── about-me.md\n"
            "│   ├── my-company.md\n"
            "│   └── procurement-standards.md\n"
            "└── OUTPUTS/\n"
            "    └── rfp-logistics/\n"
            "        ├── technical-brief.md     (input from the user)\n"
            "        ├── executive-overview.md\n"
            "        ├── scope-of-work.md\n"
            "        ├── evaluation-criteria.md\n"
            "        ├── commercial-terms.md\n"
            "        └── supplier-response-template.md"
        ),
        "behind_scenes": [
            "Read every file in ABOUT ME/ for context: who you are, your company's strategy, your scoring weights.",
            "Read the technical-brief.md you saved before the session.",
            "Asked five to eight clarifying questions about contract duration, qualifications, evaluation weighting, compliance requirements, and contract structure.",
            "Used your scoring weights from procurement-standards.md as the starting point for the evaluation criteria matrix.",
            "Wrote each of the five RFP sections as a separate Markdown file so legal and category leads can review them in parallel.",
            "Saved all five outputs to OUTPUTS/rfp-logistics/. ABOUT ME/ files were not modified.",
        ],
    },
    {
        "anchor": "Use Case 2: Spend Variance Report from Oracle Data",
        "insert_after": "variance",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "└── OUTPUTS/\n"
            "    └── spend-reports/\n"
            "        ├── oracle-q3-export.csv      (input)\n"
            "        └── q3-variance-narrative.md (output)"
        ),
        "behind_scenes": [
            "Read OUTPUTS/spend-reports/oracle-q3-export.csv. Detected the Oracle export columns (commodity code, supplier, planned, actual, variance).",
            "Calculated variance percentages and ranked commodities by absolute variance, descending.",
            "Identified the top five over-budget commodities and the top three under-budget commodities.",
            "Wrote a CFO-style narrative leading with the headline variance figure for the quarter.",
            "Quoted specific suppliers and commodity codes for each material variance.",
            "Suggested three concrete management actions tied to the worst overruns.",
            "Saved the narrative to OUTPUTS/spend-reports/q3-variance-narrative.md alongside the source CSV.",
        ],
    },
    {
        "anchor": "Use Case 3: Contract Clause Analysis",
        "insert_after": "deviation",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   └── procurement-standards.md\n"
            "└── OUTPUTS/\n"
            "    └── contract-reviews/\n"
            "        ├── <supplier>-msa.pdf       (input)\n"
            "        └── <supplier>-review.md    (output)"
        ),
        "behind_scenes": [
            "Read the source MSA PDF cover to cover. Noted clause numbering and page references.",
            "Loaded ABOUT ME/procurement-standards.md to know your company's baseline terms.",
            "Compared each clause against the matching baseline. Flagged every deviation.",
            "Classified each deviation as 'High', 'Medium', or 'Low' risk based on the standard's severity guidance.",
            "Suggested a recommended fix for each deviation (your standard wording, or a redline that meets in the middle).",
            "Wrote a three-bullet executive summary for legal at the top of the review file.",
            "Saved as <supplier>-review.md alongside the source PDF.",
        ],
    },
    {
        "anchor": "Use Case 4: Supplier Scorecard and QBR Narrative",
        "insert_after": "scorecard",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   └── procurement-standards.md\n"
            "├── TEMPLATES/\n"
            "│   └── qbr_scorecard.docx\n"
            "└── OUTPUTS/\n"
            "    └── <supplier>-qbr/\n"
            "        ├── scorecard.docx\n"
            "        └── talking-points.docx"
        ),
        "behind_scenes": [
            "Read the four KPI numbers and the QBR audience from the prompt.",
            "Loaded ABOUT ME/procurement-standards.md for the RAG bands and the scoring weights.",
            "Filled TEMPLATES/qbr_scorecard.docx with the KPIs, RAG-coloured.",
            "Drafted the talking points opening with what is going well, pivoting to the worst-performing KPI, closing with the management ask.",
            "Used the supplier's name and contract value verbatim in both documents so the reader can recognise the deal.",
            "Saved both files into a dated OUTPUTS/<supplier>-qbr/ folder.",
        ],
    },
    {
        "anchor": "Use Case 5: Pre-Negotiation Brief",
        "insert_after": "negotiation",
        "folder": (
            "Claude Cowork/\n"
            "├── ABOUT ME/\n"
            "│   ├── my-company.md\n"
            "│   └── procurement-standards.md\n"
            "└── OUTPUTS/\n"
            "    └── pre-negotiation/\n"
            "        ├── <supplier>-context.md  (input notes)\n"
            "        └── <supplier>-brief.md   (output)"
        ),
        "behind_scenes": [
            "Read OUTPUTS/pre-negotiation/<supplier>-context.md for the contract value, the open issues, and your walk-away position.",
            "Loaded ABOUT ME/my-company.md to know this year's strategic priorities for the category.",
            "Built a one-page brief with five sections: their position, our position, three points where they are exposed, three points where we are exposed, and our recommended opening.",
            "Listed two soft asks (low-cost concessions you can offer) and two hard asks (red lines).",
            "Suggested the order to raise the asks based on what they are most likely to concede first.",
            "Saved as OUTPUTS/pre-negotiation/<supplier>-brief.md.",
        ],
    },
]


# ---------------------------------------------------------------------------
# Apply augments
# ---------------------------------------------------------------------------

def apply_augments(doc, specs: Iterable[dict], label: str) -> int:
    """For each spec, find the insertion point, then inject folder + behind-scenes."""
    applied = 0
    skipped = []
    for spec in specs:
        anchor = spec["anchor"]
        after_frag = spec.get("insert_after", anchor)
        # First confirm the anchor exists.
        anchor_p = find_paragraph_by_text(doc, anchor)
        if anchor_p is None:
            skipped.append(("anchor not found", anchor))
            continue
        # Find paragraph to insert after: the first paragraph with after_frag
        # that comes after the anchor.
        target = find_paragraph_after(doc, anchor, after_frag)
        if target is None:
            # fall back: insert immediately after the anchor heading
            target = anchor_p
        # Insert in REVERSE order so each goes into the same slot:
        #   we want [target, intro_para, code_block, scenes_intro, 1., 2., 3., ...]
        # so we insert from the bottom up.
        # 1) Build the behind-scenes block first (last in document order).
        last = insert_para_after(target, "What Claude did, behind the scenes.", bold=True)
        last = insert_numbered_list_after(last, spec["behind_scenes"])
        # 2) Folder layout block before that.
        folder_intro = insert_para_after(target, "Folder layout for this scenario.", bold=True)
        insert_code_after(folder_intro, spec["folder"])
        applied += 1
    print(f"  {label}: applied {applied}, skipped {len(skipped)}")
    for reason, frag in skipped:
        print(f"    - {reason}: {frag[:80]}")
    return applied


def main():
    targets = [
        (HANDOUTS / "Claude_Code_Desktop_S2P_Training.docx", CODE_DESKTOP_AUGMENTS, "Code Desktop"),
        (HANDOUTS / "Claude_Cowork_S2P_Training.docx", COWORK_AUGMENTS, "Cowork"),
    ]
    for path, specs, label in targets:
        print(f"\n=== {label} ===")
        if not path.exists():
            print(f"  not found: {path}")
            continue
        doc = open_with_retry(path)
        apply_augments(doc, specs, label)
        # Save as a new versioned copy so the originals stay safe.
        out = path.with_name(path.stem + "_v2.docx")
        doc.save(str(out))
        print(f"  saved: {out.name}")


if __name__ == "__main__":
    main()
