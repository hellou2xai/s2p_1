"""Convert each Handouts/*.docx into an LMS-ready HTML file.

Output: Handouts/LMS/<base>_LMS.html (one per source .docx).

The converter walks the .docx body in document order and emits semantic
HTML with an inline stylesheet modeled on the existing
Course_02_The_Context_Architect_LMS.html. It preserves:

    - Cover-page block (U2xAI / PROCUREAI ACADEMY / series / title /
      subtitle / meta) inside a div.course-header.
    - Headings (Heading 1 to Heading 3 styles).
    - Bulleted and numbered lists, including ones written as plain
      paragraphs with manual "1. " or bullet prefixes.
    - Code blocks (paragraphs whose runs all use Consolas / Courier).
    - Tables, including the data-flow banner table whose cells carry
      pastel background shading (the shading is recovered from the
      cell XML and applied as inline style).
    - Folder trees (paragraphs that contain box-drawing characters).
    - Inline bold (strong), italic (em), and Consolas spans (code).
    - The "How it all connects: the data flow" banner heading and its
      colored pill table.

The script is non-destructive: existing handcrafted HTML files outside
Handouts/LMS/ (for example Course_02_The_Context_Architect_LMS.html
sitting at Handouts/ root) are untouched.

Run from the project root:

    python scripts/build_lms_html.py
"""

from __future__ import annotations

import html
import re
from pathlib import Path
from typing import List, Optional, Tuple

from docx import Document
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph


_SCRIPT_DIR = Path(__file__).resolve().parent
_PROJECT_ROOT = _SCRIPT_DIR.parent
HANDOUTS_DIR = _PROJECT_ROOT / "Handouts"
OUTPUT_DIR = HANDOUTS_DIR / "LMS"

CODE_FONTS = {"consolas", "courier new", "courier", "cascadia code", "cascadia mono", "fira code"}

BOX_CHARS = set("┌┐└┘├┤┬┴┼─│╭╮╯╰━┃┏┓┗┛")

NUMBERED_LIST_RE = re.compile(r"^\s*(\d+)[.)]\s+(.*)$", re.S)
BULLETED_LIST_RE = re.compile(r"^\s*[•\-\*]\s+(.*)$", re.S)

CALLOUT_TRIGGERS: List[Tuple[re.Pattern, str]] = [
    (re.compile(r"^\s*Why this matters\.?\s*", re.I), "why-matters"),
    (re.compile(r"^\s*What you should see\.?\s*", re.I), "callout callout-green"),
    (re.compile(r"^\s*What Claude did, behind the scenes\.?\s*", re.I), "behind-the-scenes"),
    (re.compile(r"^\s*What Claude Code did, behind the scenes\.?\s*", re.I), "behind-the-scenes"),
    (re.compile(r"^\s*What to learn from this\.?\s*", re.I), "what-to-learn"),
]

BANNER_TITLE = "How it all connects: the data flow"

STYLESHEET = """
  body {
    font-family: 'Segoe UI', Calibri, Arial, sans-serif;
    font-size: 15px;
    line-height: 1.7;
    color: #1a1a1a;
    max-width: 860px;
    margin: 0 auto;
    padding: 40px 24px;
    background: #ffffff;
  }
  .course-header {
    text-align: center;
    border-bottom: 3px solid #1a3a5c;
    padding-bottom: 24px;
    margin-bottom: 36px;
  }
  .course-header .brand {
    font-size: 13px;
    text-transform: uppercase;
    letter-spacing: 2.5px;
    color: #6b7280;
    margin-bottom: 4px;
  }
  .course-header .academy {
    font-size: 14px;
    font-weight: 600;
    text-transform: uppercase;
    letter-spacing: 2px;
    color: #1a3a5c;
    margin-bottom: 2px;
  }
  .course-header .series {
    font-size: 13px;
    color: #6b7280;
    margin-bottom: 12px;
  }
  .course-header h1 {
    font-size: 32px;
    font-weight: 700;
    color: #1a3a5c;
    margin: 8px 0 8px 0;
    letter-spacing: -0.5px;
    border: none;
    padding: 0;
  }
  .course-header .subtitle {
    font-size: 16px;
    color: #4b5563;
    font-style: italic;
    margin-bottom: 8px;
  }
  .course-header .meta {
    font-size: 14px;
    color: #6b7280;
  }
  h1 {
    font-size: 24px;
    font-weight: 700;
    color: #1a3a5c;
    border-bottom: 2px solid #e5e7eb;
    padding-bottom: 8px;
    margin-top: 48px;
    margin-bottom: 16px;
  }
  h2 {
    font-size: 19px;
    font-weight: 600;
    color: #1e40af;
    margin-top: 32px;
    margin-bottom: 12px;
  }
  h3 {
    font-size: 16px;
    font-weight: 600;
    color: #374151;
    margin-top: 24px;
    margin-bottom: 8px;
  }
  p {
    margin: 0 0 14px 0;
  }
  ul, ol {
    margin: 0 0 14px 0;
    padding-left: 24px;
  }
  li {
    margin-bottom: 6px;
  }
  pre {
    background: #1e293b;
    color: #e2e8f0;
    padding: 16px 20px;
    border-radius: 8px;
    font-family: 'Cascadia Code', 'Fira Code', Consolas, 'Courier New', monospace;
    font-size: 13.5px;
    line-height: 1.6;
    overflow-x: auto;
    margin: 12px 0 16px 0;
    white-space: pre-wrap;
    word-wrap: break-word;
  }
  code {
    background: #f1f5f9;
    color: #0f172a;
    padding: 2px 6px;
    border-radius: 4px;
    font-family: 'Cascadia Code', 'Fira Code', Consolas, 'Courier New', monospace;
    font-size: 13.5px;
  }
  pre code {
    background: none;
    color: inherit;
    padding: 0;
  }
  table {
    width: 100%;
    border-collapse: collapse;
    margin: 16px 0 20px 0;
    font-size: 14px;
  }
  thead th {
    background: #1a3a5c;
    color: #ffffff;
    font-weight: 600;
    text-align: left;
    padding: 10px 14px;
  }
  tbody td {
    border: 1px solid #d1d5db;
    padding: 10px 14px;
    vertical-align: top;
  }
  tbody tr:nth-child(even) {
    background: #f8fafc;
  }
  table.banner {
    margin: 4px 0 18px 0;
    table-layout: auto;
    width: auto;
  }
  table.banner td {
    border: none;
    padding: 8px 14px;
    text-align: center;
    vertical-align: middle;
    font-size: 13px;
    border-radius: 6px;
  }
  table.banner td.arrow {
    background: transparent;
    color: #6b7280;
    font-weight: 600;
  }
  table.banner td.pill {
    font-weight: 700;
  }
  .banner-title {
    font-weight: 700;
    color: #1a3a5c;
    margin-top: 6px;
    margin-bottom: 4px;
    font-size: 14px;
  }
  .callout {
    background: #eff6ff;
    border-left: 4px solid #2563eb;
    padding: 14px 18px;
    border-radius: 0 6px 6px 0;
    margin: 16px 0;
    font-size: 14px;
  }
  .callout-green {
    background: #f0fdf4;
    border-left: 4px solid #16a34a;
  }
  .callout-amber {
    background: #fffbeb;
    border-left: 4px solid #d97706;
  }
  .behind-the-scenes {
    background: #faf5ff;
    border-left: 4px solid #7c3aed;
    padding: 14px 18px;
    border-radius: 0 6px 6px 0;
    margin: 12px 0 16px 0;
  }
  .behind-the-scenes strong {
    color: #7c3aed;
  }
  .why-matters {
    background: #fefce8;
    border-left: 4px solid #ca8a04;
    padding: 12px 16px;
    border-radius: 0 6px 6px 0;
    margin: 10px 0 16px 0;
    font-size: 14px;
  }
  .what-to-learn {
    background: #ecfdf5;
    border-left: 4px solid #059669;
    padding: 12px 16px;
    border-radius: 0 6px 6px 0;
    margin: 10px 0 16px 0;
    font-size: 14px;
  }
  .folder-tree {
    background: #f8fafc;
    border: 1px solid #e2e8f0;
    padding: 16px 20px;
    border-radius: 8px;
    font-family: 'Cascadia Code', 'Fira Code', Consolas, 'Courier New', monospace;
    font-size: 13px;
    line-height: 1.7;
    margin: 12px 0 16px 0;
    white-space: pre;
    overflow-x: auto;
  }
  hr {
    border: none;
    border-top: 1px solid #e5e7eb;
    margin: 32px 0;
  }
"""


# ---------------------------------------------------------------------------
# Run-level helpers
# ---------------------------------------------------------------------------

def _run_font_name(run) -> str:
    name = (run.font.name or "").lower()
    if name:
        return name
    rPr = run._element.find(qn("w:rPr"))
    if rPr is not None:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is not None:
            for attr in ("ascii", "hAnsi", "cs", "eastAsia"):
                v = rFonts.get(qn(f"w:{attr}"))
                if v:
                    return v.lower()
    return ""


def _is_code_run(run) -> bool:
    return _run_font_name(run) in CODE_FONTS


def _runs(paragraph: Paragraph):
    return list(paragraph.runs)


def _is_code_block_paragraph(paragraph: Paragraph) -> bool:
    runs = _runs(paragraph)
    if not runs:
        return False
    text_runs = [r for r in runs if r.text and r.text.strip()]
    if not text_runs:
        return False
    return all(_is_code_run(r) for r in text_runs)


def _has_box_chars(text: str) -> bool:
    return any(ch in BOX_CHARS for ch in text)


def _para_text(paragraph: Paragraph) -> str:
    return paragraph.text or ""


def _runs_to_html(runs) -> str:
    out = []
    for run in runs:
        text = run.text or ""
        if not text:
            continue
        chunk = html.escape(text)
        if _is_code_run(run):
            chunk = f"<code>{chunk}</code>"
        else:
            if run.bold:
                chunk = f"<strong>{chunk}</strong>"
            if run.italic:
                chunk = f"<em>{chunk}</em>"
        out.append(chunk)
    return "".join(out)


def _para_inline_html(paragraph: Paragraph) -> str:
    return _runs_to_html(_runs(paragraph))


# ---------------------------------------------------------------------------
# Cell shading lookup (for the data-flow banner)
# ---------------------------------------------------------------------------

def _cell_fill(cell) -> Optional[str]:
    tcPr = cell._tc.find(qn("w:tcPr"))
    if tcPr is None:
        return None
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return None
    fill = shd.get(qn("w:fill"))
    if fill and fill.lower() != "auto":
        return f"#{fill}"
    return None


# ---------------------------------------------------------------------------
# HTML builder
# ---------------------------------------------------------------------------

class HtmlBuilder:
    def __init__(self):
        self.parts: List[str] = []
        self.list_stack: List[str] = []

    def open_list(self, kind: str) -> None:
        if self.list_stack and self.list_stack[-1] == kind:
            return
        self.close_lists()
        self.parts.append(f"<{kind}>")
        self.list_stack.append(kind)

    def close_lists(self) -> None:
        while self.list_stack:
            kind = self.list_stack.pop()
            self.parts.append(f"</{kind}>")

    def add_li(self, inner_html: str) -> None:
        if not self.list_stack:
            self.open_list("ul")
        self.parts.append(f"  <li>{inner_html}</li>")

    def add(self, raw_html: str) -> None:
        self.close_lists()
        self.parts.append(raw_html)

    def build(self) -> str:
        self.close_lists()
        return "\n".join(self.parts)


# ---------------------------------------------------------------------------
# Cover-page extraction
# ---------------------------------------------------------------------------

def _build_cover(paragraphs_before_first_h1: List[Paragraph]) -> Tuple[str, int]:
    """Return (html, count_of_paragraphs_consumed)."""
    if not paragraphs_before_first_h1:
        return "", 0

    brand = academy = series = title = subtitle = ""
    meta_lines: List[str] = []

    for p in paragraphs_before_first_h1:
        text = _para_text(p).strip()
        if not text:
            continue
        style = (p.style.name or "").lower() if p.style else ""
        if style == "title":
            title = text
            continue
        if text.upper() == "U2XAI":
            brand = text
        elif text.upper() == "PROCUREAI ACADEMY":
            academy = text
        elif re.match(r"^Series\s+\d", text, re.I):
            series = text
        elif not subtitle and len(text) < 120 and not _has_box_chars(text):
            subtitle = text
        else:
            meta_lines.append(text)

    if not title:
        title_p = next(
            (p for p in paragraphs_before_first_h1
             if p.style is not None
             and (p.style.name or "").lower() == "title"),
            None,
        )
        if title_p is not None:
            title = _para_text(title_p).strip()

    if not (title or brand or academy or series or subtitle or meta_lines):
        return "", 0

    body = ['<div class="course-header">']
    if brand:
        body.append(f'  <div class="brand">{html.escape(brand)}</div>')
    if academy:
        body.append(f'  <div class="academy">{html.escape(academy)}</div>')
    if series:
        body.append(f'  <div class="series">{html.escape(series)}</div>')
    if title:
        body.append(f"  <h1>{html.escape(title)}</h1>")
    if subtitle:
        body.append(f'  <div class="subtitle">{html.escape(subtitle)}</div>')
    for m in meta_lines:
        body.append(f'  <div class="meta">{html.escape(m)}</div>')
    body.append("</div>")
    return "\n".join(body), len(paragraphs_before_first_h1)


# ---------------------------------------------------------------------------
# Document walking
# ---------------------------------------------------------------------------

def _iter_body(doc):
    """Yield ("p", Paragraph) or ("tbl", Table) tuples in document order."""
    body = doc.element.body
    paragraphs_by_el = {p._p: p for p in doc.paragraphs}
    tables_by_el = {t._tbl: t for t in doc.tables}
    for child in body.iterchildren():
        if child.tag == qn("w:p"):
            p = paragraphs_by_el.get(child)
            if p is not None:
                yield ("p", p)
        elif child.tag == qn("w:tbl"):
            t = tables_by_el.get(child)
            if t is not None:
                yield ("tbl", t)


def _emit_table(builder: HtmlBuilder, table, is_banner: bool) -> None:
    if is_banner:
        builder.add('<table class="banner">')
        for row in table.rows:
            builder.add("  <tr>")
            for cell in row.cells:
                cell_html = _runs_to_html(
                    [r for p in cell.paragraphs for r in p.runs]
                ).strip()
                if not cell_html:
                    cell_html = "&nbsp;"
                fill = _cell_fill(cell)
                stripped = re.sub(r"<[^>]+>", "", cell_html).strip()
                is_arrow = stripped in {"→", "->", "⇒"}
                cls = "arrow" if is_arrow else "pill"
                style = f' style="background:{fill};"' if fill and not is_arrow else ""
                builder.add(f'    <td class="{cls}"{style}>{cell_html}</td>')
            builder.add("  </tr>")
        builder.add("</table>")
        return

    rows = list(table.rows)
    if not rows:
        return

    first_row_bold = all(
        any(r.bold for p in cell.paragraphs for r in p.runs if r.text and r.text.strip())
        for cell in rows[0].cells
        if any(p.text.strip() for p in cell.paragraphs)
    )
    has_header = first_row_bold and len(rows) > 1

    builder.add("<table>")
    if has_header:
        builder.add("  <thead><tr>")
        for cell in rows[0].cells:
            cell_html = _runs_to_html(
                [r for p in cell.paragraphs for r in p.runs]
            ).strip()
            cell_html = re.sub(r"</?(strong|em)>", "", cell_html)
            builder.add(f"    <th>{cell_html}</th>")
        builder.add("  </tr></thead>")
        body_rows = rows[1:]
    else:
        body_rows = rows
    builder.add("  <tbody>")
    for row in body_rows:
        builder.add("    <tr>")
        for cell in row.cells:
            paragraphs_html = []
            for p in cell.paragraphs:
                inner = _para_inline_html(p)
                if inner.strip():
                    paragraphs_html.append(f"<p>{inner}</p>")
            if not paragraphs_html:
                paragraphs_html.append("&nbsp;")
            builder.add(f"      <td>{''.join(paragraphs_html)}</td>")
        builder.add("    </tr>")
    builder.add("  </tbody>")
    builder.add("</table>")


def _classify_list(paragraph: Paragraph) -> Optional[Tuple[str, str]]:
    """Return (kind, inner_html) where kind is 'ul' or 'ol', or None."""
    style = (paragraph.style.name or "").lower() if paragraph.style else ""
    if "bullet" in style:
        return ("ul", _para_inline_html(paragraph))
    if "number" in style:
        return ("ol", _para_inline_html(paragraph))

    text = _para_text(paragraph)
    if not text.strip():
        return None

    indent = paragraph.paragraph_format.left_indent
    indent_pts = indent.pt if indent is not None else 0

    m = NUMBERED_LIST_RE.match(text)
    if m and indent_pts >= 18:
        # Strip the leading "N." from the inline html.
        inner = _para_inline_html(paragraph)
        inner = re.sub(r"^\s*\d+[.)]\s+", "", inner, count=1)
        return ("ol", inner)

    m = BULLETED_LIST_RE.match(text)
    if m:
        inner = _para_inline_html(paragraph)
        inner = re.sub(r"^\s*[•\-\*]\s+", "", inner, count=1)
        return ("ul", inner)

    return None


def _maybe_callout(paragraph: Paragraph) -> Optional[str]:
    """If the paragraph opens with a recognized callout label, return the
    full callout HTML. Otherwise return None.
    """
    text = _para_text(paragraph)
    for pattern, css_class in CALLOUT_TRIGGERS:
        if pattern.match(text):
            inner = _para_inline_html(paragraph)
            return f'<div class="{css_class}">{inner}</div>'
    return None


def _emit_paragraph(builder: HtmlBuilder, paragraph: Paragraph) -> None:
    text = _para_text(paragraph)
    if not text.strip():
        builder.close_lists()
        return

    style = (paragraph.style.name or "").lower() if paragraph.style else ""

    # Headings
    if style.startswith("heading "):
        level = style.split()[-1]
        try:
            n = int(level)
        except ValueError:
            n = 2
        n = max(1, min(n, 4))
        # Handout author often used Heading 1 for top-level numbered sections
        # ("1. How to use this handout"). Map to <h1>; Heading 2 -> <h2>; etc.
        builder.add(f"<h{n}>{html.escape(text.strip())}</h{n}>")
        return

    # Folder tree
    if _has_box_chars(text):
        # Multi-line content is preserved verbatim (escaped).
        builder.add(f'<div class="folder-tree">{html.escape(text)}</div>')
        return

    # Code block
    if _is_code_block_paragraph(paragraph):
        builder.add(
            f"<pre><code>{html.escape(text)}</code></pre>"
        )
        return

    # List item
    list_info = _classify_list(paragraph)
    if list_info is not None:
        kind, inner = list_info
        builder.open_list(kind)
        builder.add_li(inner)
        return

    # Callout
    callout = _maybe_callout(paragraph)
    if callout:
        builder.add(callout)
        return

    builder.add(f"<p>{_para_inline_html(paragraph)}</p>")


def _is_banner_heading(paragraph: Paragraph) -> bool:
    text = _para_text(paragraph).strip()
    return text == BANNER_TITLE


def convert(docx_path: Path) -> str:
    doc = Document(str(docx_path))

    # Cover page = the run of paragraphs before the first Heading 1, but
    # we walk the body in document order so we can also pick up the
    # banner table if it sits between the cover and the first heading.
    body_items = list(_iter_body(doc))

    # Find the index of the first Heading 1 paragraph.
    first_h1_idx = None
    for i, (kind, item) in enumerate(body_items):
        if kind == "p":
            style = (item.style.name or "").lower() if item.style else ""
            if style == "heading 1":
                first_h1_idx = i
                break

    cover_paragraphs: List[Paragraph] = []
    middle_items: List[Tuple[str, object]] = []
    if first_h1_idx is None:
        cover_paragraphs = []
        middle_items = body_items
    else:
        for kind, item in body_items[:first_h1_idx]:
            if kind == "p":
                # Banner heading + spacer + table get pulled into middle so
                # they render after the cover header div.
                if _is_banner_heading(item) or _para_text(item).strip() == "":
                    middle_items.append((kind, item))
                else:
                    cover_paragraphs.append(item)
            else:
                middle_items.append((kind, item))
        middle_items.extend(body_items[first_h1_idx:])

    builder = HtmlBuilder()
    cover_html, _ = _build_cover(cover_paragraphs)
    if cover_html:
        builder.add(cover_html)

    pending_banner_heading = False
    for idx, (kind, item) in enumerate(middle_items):
        if kind == "p":
            if _is_banner_heading(item):
                builder.add(
                    f'<div class="banner-title">{html.escape(BANNER_TITLE)}</div>'
                )
                pending_banner_heading = True
                continue
            _emit_paragraph(builder, item)
        else:  # tbl
            is_banner = pending_banner_heading
            pending_banner_heading = False
            _emit_table(builder, item, is_banner=is_banner)

    body_html = builder.build()

    title_text = ""
    if cover_paragraphs:
        for p in cover_paragraphs:
            style = (p.style.name or "").lower() if p.style else ""
            if style == "title":
                title_text = _para_text(p).strip()
                break
    if not title_text:
        title_text = docx_path.stem.replace("_", " ")

    return f"""<!DOCTYPE html>
<html lang=\"en\">
<head>
<meta charset=\"UTF-8\">
<meta name=\"viewport\" content=\"width=device-width, initial-scale=1.0\">
<title>{html.escape(title_text)}</title>
<style>{STYLESHEET}</style>
</head>
<body>
{body_html}
</body>
</html>
"""


# ---------------------------------------------------------------------------
# Driver
# ---------------------------------------------------------------------------

def _output_path(docx_path: Path) -> Path:
    base = docx_path.stem
    base = re.sub(r"_Handout(_v\d+)?$", lambda m: (m.group(1) or ""), base)
    if not base.endswith("_LMS"):
        base = f"{base}_LMS"
    rel_parent = docx_path.parent.relative_to(HANDOUTS_DIR)
    out_dir = OUTPUT_DIR / rel_parent
    return out_dir / f"{base}.html"


def main() -> None:
    if not HANDOUTS_DIR.exists():
        print(f"Handouts directory not found: {HANDOUTS_DIR}")
        return

    files = sorted(HANDOUTS_DIR.rglob("*.docx"))
    files = [f for f in files if " - Copy" not in f.name]
    files = [f for f in files if not f.name.startswith("~$")]

    if not files:
        print("No .docx handouts found.")
        return

    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    for f in files:
        try:
            html_text = convert(f)
            out_path = _output_path(f)
            out_path.parent.mkdir(parents=True, exist_ok=True)
            out_path.write_text(html_text, encoding="utf-8")
            status = f"-> {out_path.relative_to(_PROJECT_ROOT)}"
        except Exception as exc:
            status = f"ERROR: {exc}"
        rel = f.relative_to(_PROJECT_ROOT)
        print(f"{rel}: {status}")


if __name__ == "__main__":
    main()
