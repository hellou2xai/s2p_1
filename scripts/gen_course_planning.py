"""
Generate ProcureAI_Course_Planning.xlsx inside Detailed Course Content/.
Run: python scripts/gen_course_planning.py
"""

import os
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
HANDOUTS_DIR = PROJECT_DIR / "Handouts"
OUTPUT_PATH = PROJECT_DIR / "Detailed Course Content" / "ProcureAI_Course_Planning.xlsx"

# ── Colors ──
NAVY = "1A3C6E"
WHITE = "FFFFFF"
LIGHT_BLUE = "D6E4F0"
LIGHT_GREEN = "E2EFDA"
LIGHT_YELLOW = "FFF2CC"
LIGHT_ORANGE = "FCE4D6"
LIGHT_PURPLE = "E8D5F5"
LIGHT_GRAY = "F2F2F2"
GREEN_PASS = "C6EFCE"

HEADER_FONT = Font(name="Calibri", size=12, bold=True, color=WHITE)
CHAPTER_FONT = Font(name="Calibri", size=11, bold=True, color=NAVY)
COURSE_FONT = Font(name="Calibri", size=11, bold=True)
NORMAL_FONT = Font(name="Calibri", size=10)
SMALL_FONT = Font(name="Calibri", size=9, italic=True, color="666666")
BOLD_FONT = Font(name="Calibri", size=11, bold=True)

HEADER_FILL = PatternFill(start_color=NAVY, end_color=NAVY, fill_type="solid")
COURSE_FILL = PatternFill(start_color=LIGHT_GRAY, end_color=LIGHT_GRAY, fill_type="solid")
GREEN_FILL = PatternFill(start_color=GREEN_PASS, end_color=GREEN_PASS, fill_type="solid")

CHAPTER_FILLS = {
    1: PatternFill(start_color=LIGHT_BLUE, end_color=LIGHT_BLUE, fill_type="solid"),
    2: PatternFill(start_color=LIGHT_GREEN, end_color=LIGHT_GREEN, fill_type="solid"),
    3: PatternFill(start_color=LIGHT_YELLOW, end_color=LIGHT_YELLOW, fill_type="solid"),
    4: PatternFill(start_color=LIGHT_ORANGE, end_color=LIGHT_ORANGE, fill_type="solid"),
    5: PatternFill(start_color=LIGHT_PURPLE, end_color=LIGHT_PURPLE, fill_type="solid"),
}

THIN_BORDER = Border(
    left=Side(style="thin", color="CCCCCC"),
    right=Side(style="thin", color="CCCCCC"),
    top=Side(style="thin", color="CCCCCC"),
    bottom=Side(style="thin", color="CCCCCC"),
)

# ── Course taxonomy ──
CHAPTERS = [
    {
        "chapter": 1,
        "title": "S2P Foundation: Setting Up Claude Code for Procurement",
        "summary": "Build the core infrastructure: context files, reusable skills, slash commands, and the mental model for working with Claude Code in a source-to-pay environment.",
        "prereq": "None (start here)",
        "courses": [
            {"num": 2, "title": "S2P with Claude Code: The Context Architect", "tagline": "How to give Claude the right background for procurement work.", "capability": "CLAUDE.md stacking (global, project, folder-level context files)", "persona": "Anwar", "focus": "Context and Configuration",
             "lessons": [("Lesson 1", "How Claude finds your CLAUDE.md files"), ("Lesson 2", "Writing the global CLAUDE.md"), ("Lesson 3", "Writing the three category files"), ("Lesson 4", "Checking the files load correctly"), ("Lesson 5", "Pointing at data files instead of repeating them"), ("Lesson 6", "One prompt, three briefs")]},
            {"num": 3, "title": "S2P with Claude Code: The Skill Builder", "tagline": "Reusable methodologies for sourcing events, captured as SKILL.md files.", "capability": "SKILL.md files for repeatable procurement methodologies", "persona": "Anwar", "focus": "Methodology and Skills",
             "lessons": [("Lesson 1", "What a SKILL.md is and is not"), ("Lesson 2", "Anatomy of a well-written skill"), ("Lesson 3", "Writing the rfp-builder skill"), ("Lesson 4", "Writing the bid-scorer skill"), ("Lesson 5", "Skill composition: chaining skills together"), ("Lesson 6", "Skill versioning: updating without breaking")]},
            {"num": 4, "title": "S2P with Claude Code: The Command Engineer", "tagline": "Repeatable slash commands for monthly procurement reporting.", "capability": "Custom slash commands with $ARGUMENTS for recurring reports", "persona": "Priya", "focus": "Reporting and Commands",
             "lessons": [("Lesson 1", "Anatomy of a custom slash command"), ("Lesson 2", "Writing /spend-analyze"), ("Lesson 3", "Writing /anomaly-detect"), ("Lesson 4", "Parameterized commands: /scorecard-refresh and /contract-sweep"), ("Lesson 5", "Chaining commands: /rfp-launch and /savings-update"), ("Lesson 6", "Team command distribution")]},
        ],
    },
    {
        "chapter": 2,
        "title": "S2P Architecture: Multi-Agent, Hooks, and Automation",
        "summary": "Scale beyond single sessions. Orchestrate parallel sub-agents, add quality gates with hooks, automate nightly pipelines, and build persistent projects for source-to-pay workflows.",
        "prereq": "Chapter 1 (S2P Foundation)",
        "courses": [
            {"num": 5, "title": "S2P with Claude Code: The Orchestrator", "tagline": "Multi-agent supplier portfolio scoring with parallel sub-agents.", "capability": "Agent tool for parallel sub-agent orchestration", "persona": "Marcus Chen", "focus": "Multi-Agent Architecture",
             "lessons": [("Lesson 1", "When sub-agents are the right design"), ("Lesson 2", "The Task tool"), ("Lesson 3", "Designing the orchestrator"), ("Lesson 4", "Designing the worker: narrow scope and structured JSON output"), ("Lesson 5", "State handoff: JSON outputs that aggregate without re-reading"), ("Lesson 6", "Failure handling: detecting malformed output and recovering")]},
            {"num": 6, "title": "S2P with Claude Code: The Guardian", "tagline": "PreToolUse and PostToolUse hooks for contract review quality gates.", "capability": "Hook scripts that validate, audit, and alert automatically", "persona": "Rachel", "focus": "Quality Gates and Hooks",
             "lessons": [("Lesson 1", "The four hook types"), ("Lesson 2", "Writing the validation hook"), ("Lesson 3", "Writing the audit hook"), ("Lesson 4", "The risk alert hook"), ("Lesson 5", "Hook error handling"), ("Lesson 6", "Integration test")]},
            {"num": 7, "title": "S2P with Claude Code: The Pipeline Automator", "tagline": "Stop hooks and notification routing for automated spend monitoring.", "capability": "Stop hooks, tiered notifications, Slack webhook integration", "persona": "Marcus", "focus": "Automation and Notifications",
             "lessons": [("Lesson 1", "Stop hooks"), ("Lesson 2", "Notification hooks"), ("Lesson 3", "Tiered notification router"), ("Lesson 4", "Slack webhook integration"), ("Lesson 5", "Complete nightly stack"), ("Lesson 6", "End-to-end test")]},
            {"num": 8, "title": "S2P with Claude Code: The Project Architect", "tagline": "Cross-session memory for savings program management.", "capability": "Persistent state files, decisions log, project-level settings", "persona": "Dana", "focus": "Project Memory and State",
             "lessons": [("Lesson 1", "What a Claude Code project is"), ("Lesson 2", "Writing a project CLAUDE.md for a multi-initiative program"), ("Lesson 3", "Persistent state files"), ("Lesson 4", "The decisions log pattern"), ("Lesson 5", "Project settings and slash commands"), ("Lesson 6", "Team project sharing")]},
        ],
    },
    {
        "chapter": 3,
        "title": "S2P Integration and Composition",
        "summary": "Connect Claude Code to live source-to-pay systems, compose sub-agents with hooks and commands into full procurement workflows, and deploy to teams.",
        "prereq": "Chapters 1 and 2",
        "courses": [
            {"num": 9, "title": "S2P with Claude Code: The Integration Architect", "tagline": "MCP servers connecting Claude Code to live procurement databases.", "capability": "MCP protocol, FastMCP, tools vs. resources, security patterns", "persona": "David", "focus": "System Integration (MCP)",
             "lessons": [("Lesson 1", "The MCP protocol"), ("Lesson 2", "Connecting to an MCP server"), ("Lesson 3", "Anatomy of an MCP server"), ("Lesson 4", "Building the procurement MCP server"), ("Lesson 5", "Security patterns for MCP servers"), ("Lesson 6", "MCP plus skills for live spend analysis")]},
            {"num": 10, "title": "S2P with Claude Code: The Negotiation Intelligence System", "tagline": "Sub-agents, hooks, and commands for contract negotiation.", "capability": "Composing sub-agents + hooks + commands into one system", "persona": "Maria Ruiz", "focus": "System Composition",
             "lessons": [("Lesson 1", "System design mapping"), ("Lesson 2", "Intelligence gathering with sub-agents"), ("Lesson 3", "Deviation costing with a PostToolUse hook"), ("Lesson 4", "Pre-negotiation brief with a PreToolUse block"), ("Lesson 5", "Counter-proposal generation"), ("Lesson 6", "Post-close capture")]},
            {"num": 11, "title": "S2P with Claude Code: The Category Management System", "tagline": "Multi-category orchestration with self-correcting feedback loops.", "capability": "Self-correcting feedback loops, retry limits, consolidated briefings", "persona": "Elena Rodriguez", "focus": "Orchestration Patterns",
             "lessons": [("Lesson 1", "The multi-category orchestrator"), ("Lesson 2", "Specialized sub-agents: one agent per category"), ("Lesson 3", "Self-correcting feedback loop"), ("Lesson 4", "Retry limits and human escalation"), ("Lesson 5", "Consolidated briefing: assembling sub-agent outputs"), ("Lesson 6", "Performance and cost awareness: token usage and budgets")]},
            {"num": 12, "title": "S2P with Claude Code: The Team Deployment Architect", "tagline": "Deploying Claude Code to a procurement team with governance.", "capability": "Shared context, audit hooks, review gates, onboarding scripts", "persona": "Kevin Wright", "focus": "Team Deployment",
             "lessons": [("Lesson 1", "Shared vs. personal boundary"), ("Lesson 2", "Audit hook for teams"), ("Lesson 3", "The review gate hook"), ("Lesson 4", "The onboarding script"), ("Lesson 5", "Governance in practice"), ("Lesson 6", "Measuring deployment success")]},
        ],
    },
    {
        "chapter": 4,
        "title": "S2P Domain Applications: Upstream Procurement",
        "summary": "Apply Claude Code to upstream source-to-pay processes: contract intelligence, supplier lifecycle, sourcing events, and market analysis.",
        "prereq": "Chapters 1 and 2 (Chapter 3 recommended)",
        "courses": [
            {"num": 13, "title": "S2P with Claude Code: Contract Intelligence", "tagline": "Extracting terms from 20 contracts, mapping obligations, building a renewal calendar.", "capability": "Contract extraction, obligation mapping, renewal tracking", "persona": "Lisa", "focus": "Contract Management",
             "lessons": [("Lesson 1", "Contract data extraction"), ("Lesson 2", "Obligation mapping"), ("Lesson 3", "Building the contract register"), ("Lesson 4", "Renewal calendar"), ("Lesson 5", "Intake trigger"), ("Lesson 6", "Contract drafting from term sheets")]},
            {"num": 14, "title": "S2P with Claude Code: Supplier Lifecycle Management", "tagline": "Tracking 30 suppliers through onboarding, active, at-risk, and exit stages.", "capability": "Lifecycle stages, state transitions, corrective action plans", "persona": "Sarah Chen", "focus": "Supplier Management",
             "lessons": [("Lesson 1", "Supplier segmentation and lifecycle stages"), ("Lesson 2", "Onboarding automation"), ("Lesson 3", "Strategic development plans"), ("Lesson 4", "At-risk detection and corrective action"), ("Lesson 5", "Exit management"), ("Lesson 6", "Relationship state persistence")]},
            {"num": 16, "title": "S2P with Claude Code: Sourcing Sprint", "tagline": "Running a competitive sourcing event from category analysis through award.", "capability": "RFP generation, bid scoring, award recommendation", "persona": "Tom Bradshaw", "focus": "Strategic Sourcing",
             "lessons": [("Lesson 1", "Category context ingestion"), ("Lesson 2", "RFP package generation"), ("Lesson 3", "Supplier response template"), ("Lesson 4", "Bid processing and scoring"), ("Lesson 5", "Award recommendation")]},
            {"num": 17, "title": "S2P with Claude Code: Market Intelligence", "tagline": "Commodity tracking, demand consolidation, and market-linked sourcing.", "capability": "Commodity price tracking, demand consolidation, make vs. buy", "persona": "Maria Rodriguez", "focus": "Market Analysis",
             "lessons": [("Lesson 1", "Market intelligence ingestion"), ("Lesson 2", "Commodity tracker"), ("Lesson 3", "Requirements processing"), ("Lesson 4", "Demand consolidation"), ("Lesson 5", "Make vs. buy analysis"), ("Lesson 6", "Market to sourcing brief")]},
        ],
    },
    {
        "chapter": 5,
        "title": "S2P Domain Applications: Downstream and Governance",
        "summary": "Apply Claude Code to downstream source-to-pay processes and governance: P2P compliance, savings tracking, supply chain risk, ESG, and audit readiness.",
        "prereq": "Chapters 1 and 2 (Chapter 3 recommended)",
        "courses": [
            {"num": 15, "title": "S2P with Claude Code: Purchase to Pay Intelligence", "tagline": "P2P compliance screening: approval checks, PO splitting, three-way match.", "capability": "Approval authority, PO compliance, three-way match, maverick spend", "persona": "Rachel", "focus": "P2P Compliance",
             "lessons": [("Lesson 1", "P2P data architecture"), ("Lesson 2", "Requisition compliance screening"), ("Lesson 3", "PO pricing compliance"), ("Lesson 4", "Three-way match automation"), ("Lesson 5", "Maverick spend detection"), ("Lesson 6", "Payment terms optimization")]},
            {"num": 18, "title": "S2P with Claude Code: Savings Program Management", "tagline": "Tracking $12M in savings across eight initiatives with CFO reporting.", "capability": "Savings methodology, transaction matching, scenario modeling", "persona": "Angela", "focus": "Savings Tracking",
             "lessons": [("Lesson 1", "Savings methodology encoding"), ("Lesson 2", "Transaction-to-initiative matching"), ("Lesson 3", "Three-scenario modeling"), ("Lesson 4", "Writing the CFO memo"), ("Lesson 5", "Monthly savings refresh")]},
            {"num": 19, "title": "S2P with Claude Code: Supply Chain Risk", "tagline": "Risk scoring, single-source mapping, disruption modeling, board brief.", "capability": "Risk signals, concentration risk, disruption scenarios, board reporting", "persona": "James", "focus": "Risk Management",
             "lessons": [("Lesson 1", "Risk signal architecture"), ("Lesson 2", "Concentration risk"), ("Lesson 3", "Single-source exposure mapping"), ("Lesson 4", "Disruption scenario modeling"), ("Lesson 5", "Mitigation planning"), ("Lesson 6", "The board resilience brief")]},
            {"num": 20, "title": "S2P with Claude Code: ESG and Sustainable Procurement", "tagline": "ESG scoring, Scope 3 estimation, supplier action plans, portfolio dashboard.", "capability": "ESG framework, carbon estimation, action plans, board dashboard", "persona": "Sarah", "focus": "ESG and Sustainability",
             "lessons": [("Lesson 1", "ESG framework design"), ("Lesson 2", "Assessment processing"), ("Lesson 3", "Scope 3 carbon estimation"), ("Lesson 4", "ESG scoring and red flag detection"), ("Lesson 5", "Supplier action plans"), ("Lesson 6", "Portfolio dashboard and board reporting")]},
            {"num": 21, "title": "S2P with Claude Code: Compliance, Policy, and Audit Readiness", "tagline": "Policy encoding, compliance checking, append-only ledger, audit packaging.", "capability": "Policy rules, compliance checks, compliance ledger, audit evidence", "persona": "Robert", "focus": "Compliance and Audit",
             "lessons": [("Lesson 1", "Encoding procurement policy as testable rules"), ("Lesson 2", "Approval authority compliance"), ("Lesson 3", "Preferred supplier compliance"), ("Lesson 4", "Documentation completeness check"), ("Lesson 5", "The compliance ledger"), ("Lesson 6", "Audit package assembly")]},
        ],
    },
]


def set_col_widths(ws, widths):
    for col_letter, w in widths.items():
        ws.column_dimensions[col_letter].width = w


def write_header_row(ws, row, headers):
    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = THIN_BORDER


def styled_cell(ws, row, col, value, font=NORMAL_FONT, fill=None, align=None, wrap=False):
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = font
    if fill:
        cell.fill = fill
    cell.alignment = align or Alignment(vertical="top", wrap_text=wrap)
    cell.border = THIN_BORDER
    return cell


def get_handout_stats(course_num):
    """Read word count, paragraph count, and section count from a handout .docx."""
    from docx import Document as DocxDoc
    pattern = f"Course_{course_num:02d}_"
    matches = [f for f in os.listdir(HANDOUTS_DIR) if f.startswith(pattern) and f.endswith(".docx")]
    for m in matches:
        try:
            doc = DocxDoc(str(HANDOUTS_DIR / m))
            words = sum(len(p.text.split()) for p in doc.paragraphs)
            paras = len(doc.paragraphs)
            sections = len([p for p in doc.paragraphs if p.style.name == "Heading 1"])
            return words, paras, sections
        except Exception:
            pass
    return 0, 0, 0


def build_sheet_1_overview(wb):
    ws = wb.active
    ws.title = "Course Planning"
    ws.sheet_properties.tabColor = NAVY
    set_col_widths(ws, {"A": 12, "B": 12, "C": 52, "D": 55, "E": 42, "F": 10, "G": 18})

    ws.merge_cells("A1:G1")
    ws["A1"] = "ProcureAI Academy: Source-to-Pay Foundation Series with Claude Code"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:G2")
    ws["A2"] = "20 courses (02 to 21) | 118 lessons | 5 chapters | S2P Foundation Series with Claude Code"
    ws["A2"].font = SMALL_FONT
    ws["A2"].alignment = Alignment(horizontal="center")

    write_header_row(ws, 4, ["Chapter", "Course #", "Course Title", "Summary / Tagline", "Key Claude Code Capability", "Lessons", "Persona"])

    row = 5
    for ch in CHAPTERS:
        ch_fill = CHAPTER_FILLS[ch["chapter"]]
        ch_start = row

        for ci, c in enumerate(ch["courses"]):
            a = styled_cell(ws, row, 1, f"Ch {ch['chapter']}: {ch['title']}" if ci == 0 else "", CHAPTER_FONT, ch_fill, wrap=True)
            styled_cell(ws, row, 2, f"Course {c['num']:02d}", COURSE_FONT, COURSE_FILL, Alignment(horizontal="center", vertical="top"))
            styled_cell(ws, row, 3, c["title"], COURSE_FONT, wrap=True)
            styled_cell(ws, row, 4, c["tagline"], NORMAL_FONT, wrap=True)
            styled_cell(ws, row, 5, c["capability"], NORMAL_FONT, wrap=True)
            styled_cell(ws, row, 6, len(c["lessons"]), NORMAL_FONT, align=Alignment(horizontal="center", vertical="top"))
            styled_cell(ws, row, 7, c["persona"], NORMAL_FONT, align=Alignment(horizontal="center", vertical="top"))
            row += 1

        if len(ch["courses"]) > 1:
            ws.merge_cells(start_row=ch_start, start_column=1, end_row=row - 1, end_column=1)

        styled_cell(ws, row, 1, "", fill=ch_fill)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=7)
        styled_cell(ws, row, 2, f"Chapter summary: {ch['summary']}", SMALL_FONT, ch_fill, wrap=True)
        row += 1

    ws.freeze_panes = "A5"


def build_sheet_2_lessons(wb):
    ws = wb.create_sheet("Lesson Breakdown")
    ws.sheet_properties.tabColor = "2C5F8A"
    set_col_widths(ws, {"A": 10, "B": 12, "C": 52, "D": 12, "E": 58, "F": 22})

    ws.merge_cells("A1:F1")
    ws["A1"] = "S2P Foundation Series with Claude Code: All 118 Lessons Across 20 Courses"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    write_header_row(ws, 3, ["Chapter", "Course #", "Course Title", "Lesson #", "Lesson Title", "Focus Area"])

    row = 4
    for ch in CHAPTERS:
        ch_fill = CHAPTER_FILLS[ch["chapter"]]
        for c in ch["courses"]:
            for li, (lnum, ltitle) in enumerate(c["lessons"]):
                styled_cell(ws, row, 1, f"Ch {ch['chapter']}", NORMAL_FONT, ch_fill, Alignment(horizontal="center"))
                styled_cell(ws, row, 2, f"Course {c['num']:02d}", NORMAL_FONT, COURSE_FILL if li == 0 else None, Alignment(horizontal="center"))
                styled_cell(ws, row, 3, c["title"] if li == 0 else "", COURSE_FONT if li == 0 else NORMAL_FONT, wrap=True)
                styled_cell(ws, row, 4, lnum, NORMAL_FONT, align=Alignment(horizontal="center"))
                styled_cell(ws, row, 5, ltitle, NORMAL_FONT, wrap=True)
                styled_cell(ws, row, 6, c["focus"], NORMAL_FONT, align=Alignment(horizontal="center", wrap_text=True))
                row += 1

    ws.freeze_panes = "A4"


def build_sheet_3_chapters(wb):
    ws = wb.create_sheet("Chapter Summary")
    ws.sheet_properties.tabColor = "4472C4"
    set_col_widths(ws, {"A": 10, "B": 48, "C": 65, "D": 10, "E": 10, "F": 32})

    ws.merge_cells("A1:F1")
    ws["A1"] = "S2P Foundation Series with Claude Code: Chapter Summary and Learning Progression"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    write_header_row(ws, 3, ["Chapter", "Chapter Title", "Summary", "Courses", "Lessons", "Prerequisites"])

    for ch in CHAPTERS:
        r = 3 + ch["chapter"]
        ch_fill = CHAPTER_FILLS[ch["chapter"]]
        total_lessons = sum(len(c["lessons"]) for c in ch["courses"])

        styled_cell(ws, r, 1, f"Chapter {ch['chapter']}", CHAPTER_FONT, ch_fill, Alignment(horizontal="center", vertical="top"))
        styled_cell(ws, r, 2, ch["title"], COURSE_FONT, ch_fill, wrap=True)
        styled_cell(ws, r, 3, ch["summary"], NORMAL_FONT, wrap=True)
        styled_cell(ws, r, 4, len(ch["courses"]), NORMAL_FONT, align=Alignment(horizontal="center", vertical="top"))
        styled_cell(ws, r, 5, total_lessons, NORMAL_FONT, align=Alignment(horizontal="center", vertical="top"))
        styled_cell(ws, r, 6, ch["prereq"], NORMAL_FONT, wrap=True)

    tr = 9
    ws.merge_cells(f"A{tr}:B{tr}")
    styled_cell(ws, tr, 1, "TOTAL", BOLD_FONT, align=Alignment(horizontal="right"))
    ws.cell(row=tr, column=2).border = THIN_BORDER
    styled_cell(ws, tr, 3, "5 chapters covering foundation through domain applications", SMALL_FONT)
    styled_cell(ws, tr, 4, 20, BOLD_FONT, align=Alignment(horizontal="center"))
    styled_cell(ws, tr, 5, 118, BOLD_FONT, align=Alignment(horizontal="center"))
    ws.cell(row=tr, column=6).border = THIN_BORDER

    ws.freeze_panes = "A4"


def build_sheet_4_quality(wb):
    ws = wb.create_sheet("Handout Quality")
    ws.sheet_properties.tabColor = "70AD47"
    set_col_widths(ws, {"A": 12, "B": 52, "C": 12, "D": 12, "E": 12, "F": 18, "G": 18})

    ws.merge_cells("A1:G1")
    ws["A1"] = "Handout Quality Tracker"
    ws["A1"].font = Font(name="Calibri", size=14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    write_header_row(ws, 3, ["Course #", "Course Title", "Words", "Sections", "Paragraphs", "Persona", "Status"])

    data = []
    row = 4
    total_words = 0
    for ch in CHAPTERS:
        for c in ch["courses"]:
            words, paras, sections = get_handout_stats(c["num"])
            total_words += words
            status = "Complete" if words >= 5000 else ("Needs review" if words >= 2000 else "Skeleton")

            styled_cell(ws, row, 1, f"Course {c['num']:02d}", NORMAL_FONT, align=Alignment(horizontal="center"))
            styled_cell(ws, row, 2, c["title"], NORMAL_FONT, wrap=True)
            cell_w = styled_cell(ws, row, 3, words, NORMAL_FONT, align=Alignment(horizontal="center"))
            cell_w.number_format = "#,##0"
            styled_cell(ws, row, 4, sections, NORMAL_FONT, align=Alignment(horizontal="center"))
            styled_cell(ws, row, 5, paras, NORMAL_FONT, align=Alignment(horizontal="center"))
            styled_cell(ws, row, 6, c["persona"], NORMAL_FONT, align=Alignment(horizontal="center"))
            sc = styled_cell(ws, row, 7, status, NORMAL_FONT, align=Alignment(horizontal="center"))
            if status == "Complete":
                sc.fill = GREEN_FILL
            row += 1

    styled_cell(ws, row, 1, "TOTAL", BOLD_FONT, align=Alignment(horizontal="right"))
    ws.cell(row=row, column=2).border = THIN_BORDER
    tw = styled_cell(ws, row, 3, total_words, BOLD_FONT, align=Alignment(horizontal="center"))
    tw.number_format = "#,##0"
    for c in range(4, 8):
        ws.cell(row=row, column=c).border = THIN_BORDER

    ws.freeze_panes = "A4"


def main():
    wb = openpyxl.Workbook()
    build_sheet_1_overview(wb)
    build_sheet_2_lessons(wb)
    build_sheet_3_chapters(wb)
    build_sheet_4_quality(wb)
    wb.save(str(OUTPUT_PATH))
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
