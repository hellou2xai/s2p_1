"""
Generate LMS/ProcureAI_Engineering_Course_Catalog.xlsx
Run: python scripts/gen_lms_catalog.py
"""

from pathlib import Path
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_DIR = SCRIPT_DIR.parent
OUTPUT_PATH = PROJECT_DIR / "LMS" / "ProcureAI_Engineering_Course_Catalog.xlsx"

# ── Styles ──
NAVY = "1A3A5C"
WHITE = "FFFFFF"
LIGHT_BLUE = "D6E4F0"
LIGHT_GREEN = "E2EFDA"
LIGHT_YELLOW = "FFF2CC"
LIGHT_ORANGE = "FCE4D6"
LIGHT_PURPLE = "E8D5F5"

THIN = Border(
    left=Side("thin", "BBBBBB"), right=Side("thin", "BBBBBB"),
    top=Side("thin", "BBBBBB"), bottom=Side("thin", "BBBBBB"),
)

H_FONT = Font("Calibri", 12, bold=True, color=WHITE)
H_FILL = PatternFill("solid", fgColor=NAVY)
CH_FONT = Font("Calibri", 11, bold=True, color=NAVY)
BOLD = Font("Calibri", 11, bold=True)
NORM = Font("Calibri", 10.5)
SMALL = Font("Calibri", 9.5, italic=True, color="555555")
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

CH_FILLS = {
    1: PatternFill("solid", fgColor=LIGHT_BLUE),
    2: PatternFill("solid", fgColor=LIGHT_GREEN),
    3: PatternFill("solid", fgColor=LIGHT_YELLOW),
    4: PatternFill("solid", fgColor=LIGHT_ORANGE),
    5: PatternFill("solid", fgColor=LIGHT_PURPLE),
}


def hdr(ws, row, cols):
    for i, c in enumerate(cols, 1):
        x = ws.cell(row, i, c)
        x.font = H_FONT
        x.fill = H_FILL
        x.alignment = CENTER
        x.border = THIN


def cell(ws, r, c, v, font=NORM, fill=None, align=WRAP):
    x = ws.cell(r, c, v)
    x.font = font
    if fill:
        x.fill = fill
    x.alignment = align
    x.border = THIN
    return x


# ══════════════════════════════════════════════════════════════
# COURSE DATA
# ══════════════════════════════════════════════════════════════
chapters = [
    {
        "ch": 1,
        "title": "S2P Foundation: Setting Up Claude Code for Procurement",
        "what_students_learn": (
            "How to configure Claude Code so it understands your procurement "
            "role, your categories, your data files, and your writing standards. "
            "Students build reusable context files, methodology templates, and "
            "one-line commands they will use every day."
        ),
        "procurement_application": (
            "Instead of explaining your job to Claude every session, you set it "
            "up once. After this chapter, Claude knows your suppliers, your spend "
            "data, your approval rules, and your reporting style. Every prompt you "
            "type from here on starts with the right background, the right data, "
            "and the right voice."
        ),
        "courses": [
            {
                "num": 2,
                "name": "The Context Architect: Teaching Claude Your Role, Categories, and Data",
                "what_you_learn": (
                    "How to write small configuration files (CLAUDE.md) that tell "
                    "Claude who you are, what category you manage, and where your "
                    "data lives. One file per folder, stacked automatically."
                ),
                "procurement_use": (
                    "You manage three categories (direct materials, logistics, "
                    "indirect). Instead of re-explaining each one every session, "
                    "Claude reads the right background automatically based on which "
                    "folder you open. One prompt produces three correctly scoped briefs."
                ),
                "lessons": [
                    ("How Claude finds your CLAUDE.md files",
                     "Understand the automatic file-loading mechanism so you know why folder choice matters."),
                    ("Writing the global CLAUDE.md",
                     "Create a single file that tells Claude your role, your folder rules, and your writing standards for every session."),
                    ("Writing the three category files",
                     "Create one small file per category folder so Claude knows the suppliers, data files, and rules specific to that category."),
                    ("Checking the files load correctly",
                     "Run a five-question test to confirm Claude loaded the right context before you start real work."),
                    ("Pointing at data files instead of repeating them",
                     "Reference CSVs by name in your config so Claude reads live data instead of stale copy-paste."),
                    ("One prompt, three briefs",
                     "Run the same prompt from three folders and get three correctly scoped category briefs in 20 minutes."),
                ],
            },
            {
                "num": 3,
                "name": "The Skill Builder: Saving RFP and Bid-Scoring Methods as Reusable Templates",
                "what_you_learn": (
                    "How to capture a repeatable procurement methodology (like "
                    "scoring bids or building an RFP) as a reusable SKILL.md file "
                    "that Claude follows step by step."
                ),
                "procurement_use": (
                    "Your RFP process has 12 steps. Instead of typing them every "
                    "time, you save them once as a skill file. Claude follows the "
                    "same methodology for every sourcing event, so nothing gets "
                    "missed and every RFP has the same professional shape."
                ),
                "lessons": [
                    ("What a SKILL.md is and is not",
                     "Understand the difference between a skill (a repeatable method) and a one-off prompt."),
                    ("Anatomy of a well-written skill",
                     "Learn the three sections every skill needs: inputs, steps, and output format."),
                    ("Writing the rfp-builder skill",
                     "Build a skill that generates a complete RFP package from a category brief and supplier list."),
                    ("Writing the bid-scorer skill",
                     "Build a skill that reads bid responses, scores them against weighted criteria, and ranks suppliers."),
                    ("Skill composition: chaining skills together",
                     "Connect the rfp-builder output to the bid-scorer input so one command runs the full sourcing cycle."),
                    ("Skill versioning: updating without breaking",
                     "Update a skill (add a new scoring criterion) without breaking work already in progress."),
                ],
            },
            {
                "num": 4,
                "name": "The Command Engineer: One-Line Commands for Monthly Spend and Anomaly Reports",
                "what_you_learn": (
                    "How to create custom one-line commands (slash commands) for "
                    "reports you run every week or month, like spend analysis, "
                    "anomaly detection, and scorecard refresh."
                ),
                "procurement_use": (
                    "Every month you run the same five reports. Instead of "
                    "rewriting prompts each time, you type /spend-analyze or "
                    "/anomaly-detect and the report generates in 30 seconds. Your "
                    "team uses the same commands, so everyone produces consistent outputs."
                ),
                "lessons": [
                    ("Anatomy of a custom slash command",
                     "Understand how a slash command file maps to a one-line trigger you type in the terminal."),
                    ("Writing /spend-analyze",
                     "Build a command that aggregates spend by supplier across thousands of rows and flags the top movers."),
                    ("Writing /anomaly-detect",
                     "Build a command that scans invoices for duplicate payments, round-number spikes, and approval bypasses."),
                    ("Parameterized commands: /scorecard-refresh and /contract-sweep",
                     "Add parameters so one command works for any category or any date range."),
                    ("Chaining commands: /rfp-launch and /savings-update",
                     "Connect commands so one triggers the next, building a multi-step workflow from single-line triggers."),
                    ("Team command distribution",
                     "Share your commands with colleagues so the whole team runs the same reports the same way."),
                ],
            },
        ],
    },
    {
        "ch": 2,
        "title": "S2P Architecture: Multi-Agent, Hooks, and Automation",
        "what_students_learn": (
            "How to scale Claude Code beyond one-at-a-time conversations. "
            "Students learn to run parallel agents (score 50 suppliers at once), "
            "add automatic quality checks before files are saved, automate nightly "
            "pipelines, and keep Claude's memory across sessions."
        ),
        "procurement_application": (
            "Real procurement work has volume: 50 suppliers to score, 10 contracts "
            "to review, 3,000 invoices to audit. This chapter teaches you to handle "
            "that volume by running multiple Claude agents in parallel, adding "
            "guardrails that catch mistakes before they reach your files, and "
            "scheduling overnight runs so reports are ready when you arrive Monday morning."
        ),
        "courses": [
            {
                "num": 5,
                "name": "The Orchestrator: Scoring 50 Suppliers in Parallel with Sub-Agents",
                "what_you_learn": (
                    "How to run multiple Claude sub-agents in parallel, each "
                    "scoring one supplier, then combine their results into a "
                    "single portfolio scorecard."
                ),
                "procurement_use": (
                    "You need to score 50 suppliers across five dimensions. "
                    "Instead of doing them one by one (2.5 hours), you launch "
                    "parallel agents that each handle one supplier. All 50 scores "
                    "land in 15 minutes. The orchestrator combines them into one "
                    "ranked scorecard."
                ),
                "lessons": [
                    ("When sub-agents are the right design",
                     "Recognize which procurement tasks benefit from parallel agents vs. a single session."),
                    ("The Task tool",
                     "Learn the mechanism Claude Code uses to launch, monitor, and collect results from sub-agents."),
                    ("Designing the orchestrator",
                     "Build the coordinator that assigns one supplier per agent and collects results."),
                    ("Designing the worker: narrow scope and structured JSON output",
                     "Write a focused agent prompt that scores one supplier and returns structured data."),
                    ("State handoff: JSON outputs that aggregate without re-reading",
                     "Design output formats so the orchestrator merges 50 results without re-reading source files."),
                    ("Failure handling: detecting malformed output and recovering",
                     "Handle the case where one agent fails (bad data, timeout) without losing the other 49 results."),
                ],
            },
            {
                "num": 6,
                "name": "The Guardian: Auto-Checking Contracts and Outputs Before They Save",
                "what_you_learn": (
                    "How to add automatic quality gates (hooks) that run every "
                    "time Claude reads or writes a file. These catch mistakes "
                    "before they reach your output."
                ),
                "procurement_use": (
                    "Before Claude saves a contract review, a hook checks: are all "
                    "required clauses addressed? Is the risk rating present? Is the "
                    "reviewer name filled in? If anything is missing, Claude stops "
                    "and tells you what to fix. No incomplete deliverables leave your desk."
                ),
                "lessons": [
                    ("The four hook types",
                     "Understand PreToolUse, PostToolUse, Notification, and Stop hooks and when each fires."),
                    ("Writing the validation hook",
                     "Build a hook that checks every output file for required fields before Claude saves it."),
                    ("Writing the audit hook",
                     "Build a hook that logs every file Claude touches to an append-only audit trail."),
                    ("The risk alert hook",
                     "Build a hook that flags high-value contracts or at-risk suppliers the moment Claude encounters them."),
                    ("Hook error handling",
                     "Handle the case where a hook itself fails without blocking all of Claude's work."),
                    ("Integration test",
                     "Run a full contract review with all hooks active and verify the audit trail is complete."),
                ],
            },
            {
                "num": 7,
                "name": "The Pipeline Automator: Overnight Spend Scans with Slack Alerts",
                "what_you_learn": (
                    "How to schedule Claude Code to run overnight or on a timer, "
                    "with automatic notifications sent to Slack or email when "
                    "something needs attention."
                ),
                "procurement_use": (
                    "Every Monday morning, your spend anomaly report is waiting in "
                    "your inbox. Claude ran it overnight: scanned 5,000 invoices, "
                    "found 12 anomalies, sent a Slack alert for the 3 above "
                    "$50,000, and saved the full report to your output folder. You "
                    "review it with coffee instead of building it from scratch."
                ),
                "lessons": [
                    ("Stop hooks",
                     "Build a hook that halts a pipeline when a critical threshold is breached (e.g., spend anomaly above $100,000)."),
                    ("Notification hooks",
                     "Build a hook that sends a message when Claude finishes a task or finds something worth flagging."),
                    ("Tiered notification router",
                     "Route alerts by severity: informational to a log, warnings to Slack, critical to email and Slack."),
                    ("Slack webhook integration",
                     "Connect Claude Code to your team's Slack channel for real-time procurement alerts."),
                    ("Complete nightly stack",
                     "Assemble a full nightly pipeline: data refresh, anomaly scan, report generation, notification."),
                    ("End-to-end test",
                     "Run the complete pipeline on practice data and verify every step fires in order."),
                ],
            },
            {
                "num": 8,
                "name": "The Project Architect: Giving Claude Memory Across Sessions for Savings Programs",
                "what_you_learn": (
                    "How to give Claude memory across sessions so it remembers "
                    "decisions, tracks initiative status, and picks up where "
                    "you left off."
                ),
                "procurement_use": (
                    "You manage a savings program with eight initiatives running "
                    "in parallel. Claude remembers which ones are on track, which "
                    "are stalled, and what decisions were made last week. When you "
                    "start a new session, Claude already knows the context. No "
                    "re-explaining."
                ),
                "lessons": [
                    ("What a Claude Code project is",
                     "Understand the project boundary, the settings file, and what persists between sessions."),
                    ("Writing a project CLAUDE.md for a multi-initiative program",
                     "Create a project-level config for a savings program spanning eight workstreams."),
                    ("Persistent state files",
                     "Store initiative status, milestones, and blockers in files Claude reads and updates each session."),
                    ("The decisions log pattern",
                     "Keep an append-only log of decisions so Claude (and your team) can trace why choices were made."),
                    ("Project settings and slash commands",
                     "Configure project-specific commands like /savings-status and /initiative-update."),
                    ("Team project sharing",
                     "Share the project with colleagues so everyone sees the same status and the same decision history."),
                ],
            },
        ],
    },
    {
        "ch": 3,
        "title": "S2P Integration and Composition",
        "what_students_learn": (
            "How to connect Claude Code to live procurement systems (ERP, contract "
            "databases, spend cubes), compose multiple features into complete "
            "workflows, and deploy the setup to a team with governance controls."
        ),
        "procurement_application": (
            "Claude Code stops being a standalone tool and becomes part of your "
            "procurement operating system. It pulls live data from your ERP, runs "
            "a full negotiation preparation workflow (intelligence, costing, brief, "
            "counter-proposal), orchestrates category reviews across your whole "
            "portfolio, and your entire team uses it with audit trails and approval gates."
        ),
        "courses": [
            {
                "num": 9,
                "name": "The Integration Architect: Connecting Claude to Live ERP and Spend Data",
                "what_you_learn": (
                    "How to connect Claude Code to external databases and APIs "
                    "using the MCP protocol, so Claude reads live procurement "
                    "data instead of static CSVs."
                ),
                "procurement_use": (
                    "Instead of exporting spend data to a CSV every week, Claude "
                    "connects directly to your procurement database. When you ask "
                    "'show me top 10 suppliers by spend this quarter,' Claude "
                    "queries live data and returns current figures."
                ),
                "lessons": [
                    ("The MCP protocol",
                     "Understand how Claude Code talks to external systems through a standardized protocol."),
                    ("Connecting to an MCP server",
                     "Connect Claude Code to an existing MCP server and verify it can read data."),
                    ("Anatomy of an MCP server",
                     "Understand the structure of an MCP server: tools, resources, and security boundaries."),
                    ("Building the procurement MCP server",
                     "Build a server that exposes supplier data, spend data, and contract data to Claude Code."),
                    ("Security patterns for MCP servers",
                     "Implement read-only access, query limits, and data masking for sensitive fields."),
                    ("MCP plus skills for live spend analysis",
                     "Combine a live data connection with a skill file to run spend analysis on real-time data."),
                ],
            },
            {
                "num": 10,
                "name": "The Negotiation Intelligence System: Full Contract Prep from Intel to Counter-Proposal",
                "what_you_learn": (
                    "How to combine sub-agents, hooks, and commands into one "
                    "integrated system that prepares you for a contract negotiation "
                    "from start to finish."
                ),
                "procurement_use": (
                    "Before a negotiation with a $2.4M supplier, Claude runs the "
                    "full prep: gathers market intelligence, costs every deviation "
                    "from standard terms, writes the pre-negotiation brief, "
                    "generates counter-proposals, and after the deal closes, "
                    "captures the agreed terms for the contract register."
                ),
                "lessons": [
                    ("System design mapping",
                     "Map the negotiation workflow to Claude Code components: which parts are agents, hooks, and commands."),
                    ("Intelligence gathering with sub-agents",
                     "Launch parallel agents to gather market pricing, supplier history, and benchmark data."),
                    ("Deviation costing with a PostToolUse hook",
                     "Automatically calculate the cost of every non-standard clause the supplier proposes."),
                    ("Pre-negotiation brief with a PreToolUse block",
                     "Generate a one-page brief with walk-away price, target price, and key leverage points."),
                    ("Counter-proposal generation",
                     "Given supplier terms, generate three counter-proposals ranked by value to your organization."),
                    ("Post-close capture",
                     "After the deal, capture agreed terms, savings achieved, and next review date into the contract register."),
                ],
            },
            {
                "num": 11,
                "name": "The Category Management System: Quarterly Reviews Across All Categories at Once",
                "what_you_learn": (
                    "How to orchestrate Claude across multiple categories "
                    "simultaneously, with feedback loops that catch and correct "
                    "errors before they reach the final output."
                ),
                "procurement_use": (
                    "Quarterly category review: Claude scores all categories in "
                    "parallel, checks each score against last quarter's data, flags "
                    "anything that looks wrong, retries the calculation, and assembles "
                    "a consolidated briefing for your CPO. One command, all categories, "
                    "30 minutes."
                ),
                "lessons": [
                    ("The multi-category orchestrator",
                     "Build an orchestrator that launches one agent per category and collects all results."),
                    ("Specialized sub-agents: one agent per category",
                     "Write focused agents that understand direct materials differently from logistics or indirect."),
                    ("Self-correcting feedback loop",
                     "Build a check that compares outputs to expected ranges and re-runs any that look wrong."),
                    ("Retry limits and human escalation",
                     "Set a retry cap so Claude asks you for help instead of looping forever on bad data."),
                    ("Consolidated briefing: assembling sub-agent outputs",
                     "Merge all category results into one executive briefing with consistent formatting."),
                    ("Performance and cost awareness: token usage and budgets",
                     "Monitor how many tokens each agent uses and set budgets to control costs."),
                ],
            },
            {
                "num": 12,
                "name": "The Team Deployment Architect: Rolling Out Claude to Your Procurement Team with Audit Trails",
                "what_you_learn": (
                    "How to deploy your Claude Code setup to a procurement team "
                    "with shared standards, audit trails, and approval gates."
                ),
                "procurement_use": (
                    "Your team of six analysts all use the same Claude Code setup. "
                    "Every output is logged. High-value contract reviews require a "
                    "second pair of eyes before they are finalized. New team members "
                    "are productive on day one because the onboarding script sets "
                    "everything up."
                ),
                "lessons": [
                    ("Shared vs. personal boundary",
                     "Decide what context is shared (team rules, folder structure) vs. personal (individual preferences)."),
                    ("Audit hook for teams",
                     "Build a hook that logs every Claude action across all team members to a shared audit file."),
                    ("The review gate hook",
                     "Build a hook that requires human approval before Claude saves outputs above a value threshold."),
                    ("The onboarding script",
                     "Create a script that sets up a new team member's Claude Code environment in 10 minutes."),
                    ("Governance in practice",
                     "Run a realistic scenario with audit, review gate, and team commands all active."),
                    ("Measuring deployment success",
                     "Track adoption metrics: how many prompts per analyst, time saved per week, error rates."),
                ],
            },
        ],
    },
    {
        "ch": 4,
        "title": "S2P Domain Applications: Upstream Procurement",
        "what_students_learn": (
            "How to apply Claude Code to the upstream half of source-to-pay: "
            "managing contracts, tracking supplier lifecycles, running sourcing "
            "events, and analyzing commodity markets."
        ),
        "procurement_application": (
            "These are the tasks that consume most of a category manager's week. "
            "Extract terms from 20 contracts in an hour instead of a week. Track "
            "30 suppliers through lifecycle stages with automatic alerts. Run a "
            "full sourcing event from RFP to award recommendation. Monitor commodity "
            "prices and consolidate demand across business units."
        ),
        "courses": [
            {
                "num": 13,
                "name": "Contract Intelligence: Extracting Terms, Mapping Obligations, and Tracking Renewals",
                "what_you_learn": (
                    "How to extract key terms from contracts (PDF and Word), map "
                    "obligations, build a contract register, and generate renewal "
                    "alerts automatically."
                ),
                "procurement_use": (
                    "You inherited 20 contracts with no register. Claude reads all "
                    "20, extracts effective dates, expiry dates, auto-renewal "
                    "clauses, termination notice periods, and payment terms. In one "
                    "hour you have a complete register and a renewal calendar with "
                    "90-day, 60-day, and 30-day alerts."
                ),
                "lessons": [
                    ("Contract data extraction",
                     "Extract key terms (dates, values, parties, clauses) from PDF and Word contracts."),
                    ("Obligation mapping",
                     "Map obligations by party: what you owe the supplier, what they owe you, with deadlines."),
                    ("Building the contract register",
                     "Assemble extracted data into a structured register with one row per contract."),
                    ("Renewal calendar",
                     "Generate a calendar with alerts at 90, 60, and 30 days before each renewal or expiry."),
                    ("Intake trigger",
                     "Build a workflow that automatically extracts terms when a new contract lands in the inbox folder."),
                    ("Contract drafting from term sheets",
                     "Given a term sheet, generate a first-draft contract using your standard template and clauses."),
                ],
            },
            {
                "num": 14,
                "name": "Supplier Lifecycle Management: Onboarding, Risk Alerts, and Exit Planning for 30 Suppliers",
                "what_you_learn": (
                    "How to track suppliers through lifecycle stages (onboarding, "
                    "active, watch, at-risk, exit) with automatic state transitions "
                    "and corrective action plans."
                ),
                "procurement_use": (
                    "You manage 30 suppliers. Claude tracks each one's stage, "
                    "flags when a supplier moves to at-risk (late deliveries, "
                    "quality issues, financial warning), generates a corrective "
                    "action plan, and alerts you 90 days before contract end for "
                    "exit planning."
                ),
                "lessons": [
                    ("Supplier segmentation and lifecycle stages",
                     "Define lifecycle stages and the criteria for moving a supplier between them."),
                    ("Onboarding automation",
                     "Generate an onboarding checklist, track completion, and move the supplier to active when done."),
                    ("Strategic development plans",
                     "For strategic suppliers, generate a development plan with joint initiatives and KPIs."),
                    ("At-risk detection and corrective action",
                     "Automatically flag suppliers that breach KPI thresholds and generate a corrective action plan."),
                    ("Exit management",
                     "Plan and execute a supplier exit: replacement sourcing, transition timeline, knowledge transfer."),
                    ("Relationship state persistence",
                     "Store lifecycle state so Claude remembers each supplier's stage across sessions."),
                ],
            },
            {
                "num": 16,
                "name": "Sourcing Sprint: Running an RFP from Category Analysis to Award Recommendation",
                "what_you_learn": (
                    "How to run a competitive sourcing event end to end: from "
                    "category analysis through RFP, bid scoring, and award "
                    "recommendation."
                ),
                "procurement_use": (
                    "You need to re-source office supplies across three business "
                    "units. Claude analyzes the category, generates the RFP "
                    "package, creates the supplier response template, processes "
                    "8 bid responses, scores them against weighted criteria, and "
                    "writes the award recommendation with a savings estimate. "
                    "Two days instead of three weeks."
                ),
                "lessons": [
                    ("Category context ingestion",
                     "Load category data (current suppliers, spend, contracts) to set the sourcing baseline."),
                    ("RFP package generation",
                     "Generate a complete RFP: scope, requirements, evaluation criteria, timeline, and response template."),
                    ("Supplier response template",
                     "Create a structured template that suppliers fill in, making bid comparison straightforward."),
                    ("Bid processing and scoring",
                     "Read 8 bid responses (PDF), extract pricing and terms, score against weighted criteria, rank."),
                    ("Award recommendation",
                     "Write the award recommendation: recommended supplier, savings vs. incumbent, risks, and next steps."),
                ],
            },
            {
                "num": 17,
                "name": "Market Intelligence: Commodity Tracking, Demand Consolidation, and Sourcing Briefs",
                "what_you_learn": (
                    "How to track commodity prices, consolidate demand across "
                    "business units, and link market data to sourcing decisions."
                ),
                "procurement_use": (
                    "Steel prices jumped 12% this quarter. Claude tracks the "
                    "commodity index, maps it to your steel suppliers, calculates "
                    "the impact on your direct materials spend, consolidates demand "
                    "across three plants, and writes a market-linked sourcing brief "
                    "recommending when to lock in pricing."
                ),
                "lessons": [
                    ("Market intelligence ingestion",
                     "Load commodity price data, supplier market reports, and industry benchmarks into Claude's context."),
                    ("Commodity tracker",
                     "Build a tracker that monitors price movements and flags when a commodity crosses a threshold."),
                    ("Requirements processing",
                     "Consolidate purchase requirements from multiple business units into one demand view."),
                    ("Demand consolidation",
                     "Aggregate demand across plants and regions to identify volume leverage opportunities."),
                    ("Make vs. buy analysis",
                     "Compare internal production cost against supplier quotes to recommend make or buy."),
                    ("Market to sourcing brief",
                     "Write a brief connecting market movements to sourcing actions: lock in, defer, or re-source."),
                ],
            },
        ],
    },
    {
        "ch": 5,
        "title": "S2P Domain Applications: Downstream and Governance",
        "what_students_learn": (
            "How to apply Claude Code to the downstream half of source-to-pay "
            "and governance: P2P compliance, savings validation, supply chain "
            "risk, ESG reporting, and audit readiness."
        ),
        "procurement_application": (
            "These are the tasks that keep procurement accountable. Screen every "
            "requisition and invoice for compliance. Validate $12M in savings "
            "claims against actual transactions. Map single-source exposure across "
            "your supply chain. Score suppliers on ESG criteria. Package audit "
            "evidence in hours instead of weeks."
        ),
        "courses": [
            {
                "num": 15,
                "name": "Purchase to Pay Intelligence: Approval Checks, Three-Way Match, and Maverick Spend Detection",
                "what_you_learn": (
                    "How to screen requisitions, purchase orders, and invoices for "
                    "compliance: approval authority, PO splitting, three-way match, "
                    "and maverick spend."
                ),
                "procurement_use": (
                    "3,000 invoices land every month. Claude checks each one: does "
                    "the approval match the authority matrix? Is this PO split to "
                    "avoid a threshold? Does the invoice match the PO and the goods "
                    "receipt? Which invoices bypassed the PO process entirely? You "
                    "get a compliance report in 5 minutes instead of sampling 50 "
                    "invoices over two days."
                ),
                "lessons": [
                    ("P2P data architecture",
                     "Set up the data model: requisitions, POs, goods receipts, invoices, and the approval matrix."),
                    ("Requisition compliance screening",
                     "Check every requisition against the approval authority matrix and flag violations."),
                    ("PO pricing compliance",
                     "Compare PO prices against contracted rates and flag overcharges above a tolerance."),
                    ("Three-way match automation",
                     "Match invoices to POs to goods receipts and flag mismatches by quantity, price, or date."),
                    ("Maverick spend detection",
                     "Find invoices with no matching PO (off-contract spend) and quantify the maverick percentage."),
                    ("Payment terms optimization",
                     "Analyze payment timing across all invoices and calculate the value of moving to optimal terms."),
                ],
            },
            {
                "num": 18,
                "name": "Savings Program Management: Validating $12M in Claims and Writing the CFO Memo",
                "what_you_learn": (
                    "How to track a multi-initiative savings program: encode "
                    "savings methodologies, match claimed savings to actual "
                    "transactions, model scenarios, and write the CFO memo."
                ),
                "procurement_use": (
                    "Your CFO expects $12M in savings this year across eight "
                    "initiatives. Claude validates each claim against real "
                    "transaction data, models three scenarios (optimistic, base, "
                    "conservative), and writes a monthly CFO memo with the current "
                    "run-rate, at-risk initiatives, and recommended actions. No "
                    "more spreadsheet gymnastics."
                ),
                "lessons": [
                    ("Savings methodology encoding",
                     "Define how each savings type is calculated: negotiated, volume, substitution, demand reduction."),
                    ("Transaction-to-initiative matching",
                     "Match actual POs and invoices to savings initiatives to validate claimed vs. realized savings."),
                    ("Three-scenario modeling",
                     "Model optimistic, base, and conservative outcomes for each initiative based on current run-rate."),
                    ("Writing the CFO memo",
                     "Generate a one-page memo: total savings, initiative status, at-risk items, three recommendations."),
                    ("Monthly savings refresh",
                     "Build a command that re-runs the full savings calculation with updated transaction data each month."),
                ],
            },
            {
                "num": 19,
                "name": "Supply Chain Risk: Single-Source Mapping, Disruption Scenarios, and Board Briefs",
                "what_you_learn": (
                    "How to score supplier risk, map single-source exposure, "
                    "model disruption scenarios, and write a board-ready "
                    "resilience brief."
                ),
                "procurement_use": (
                    "Your board asks: 'What happens if our top steel supplier "
                    "shuts down for 60 days?' Claude maps single-source "
                    "dependencies, models the revenue impact of a 60-day "
                    "disruption ($4.2M), identifies two alternative suppliers "
                    "with 30-day qualification timelines, and writes a one-page "
                    "board brief with a mitigation plan."
                ),
                "lessons": [
                    ("Risk signal architecture",
                     "Define risk signals: financial health, geographic concentration, delivery performance, compliance status."),
                    ("Concentration risk",
                     "Calculate what percentage of spend goes to each supplier and flag dangerous concentrations."),
                    ("Single-source exposure mapping",
                     "Identify every product or service where only one supplier exists and quantify the exposure."),
                    ("Disruption scenario modeling",
                     "Model 'what if supplier X is offline for N days' with revenue impact and recovery timeline."),
                    ("Mitigation planning",
                     "For each high-risk scenario, generate a mitigation plan: alternative suppliers, safety stock, dual-sourcing."),
                    ("The board resilience brief",
                     "Write a one-page brief with risk heat map, top 3 exposures, and recommended board actions."),
                ],
            },
            {
                "num": 20,
                "name": "ESG and Sustainable Procurement: Scope 3 Estimates, Supplier Scores, and Action Plans",
                "what_you_learn": (
                    "How to score suppliers on ESG criteria, estimate Scope 3 "
                    "carbon emissions from procurement data, generate supplier "
                    "action plans, and build a portfolio dashboard."
                ),
                "procurement_use": (
                    "Your sustainability team needs Scope 3 estimates for 50 "
                    "suppliers by month-end. Claude reads each supplier's ESG "
                    "assessment, scores them on your framework, estimates carbon "
                    "emissions from spend and commodity data, flags red-flag "
                    "suppliers, and generates action plans for the bottom 10. "
                    "Dashboard and board report included."
                ),
                "lessons": [
                    ("ESG framework design",
                     "Define your ESG scoring framework: environmental, social, governance criteria with weights."),
                    ("Assessment processing",
                     "Read supplier ESG self-assessments (PDFs) and extract scores for each criterion."),
                    ("Scope 3 carbon estimation",
                     "Estimate Scope 3 emissions from procurement spend data using emission factor tables."),
                    ("ESG scoring and red flag detection",
                     "Score all suppliers, rank them, and flag those below the minimum threshold."),
                    ("Supplier action plans",
                     "Generate improvement action plans for low-scoring suppliers with deadlines and KPIs."),
                    ("Portfolio dashboard and board reporting",
                     "Build a portfolio ESG dashboard and a one-page board summary with trends and targets."),
                ],
            },
            {
                "num": 21,
                "name": "Compliance, Policy, and Audit Readiness: Encoding Rules, Scanning Transactions, and Packaging Evidence",
                "what_you_learn": (
                    "How to encode procurement policies as testable rules, run "
                    "compliance checks across all transactions, maintain an "
                    "append-only compliance ledger, and package audit evidence."
                ),
                "procurement_use": (
                    "External audit in three weeks. Claude encodes your 14 "
                    "procurement policies as testable rules, scans every "
                    "transaction from the last 12 months, logs results to a "
                    "tamper-proof ledger, and packages the evidence (test results, "
                    "supporting documents, exception approvals) into an audit-ready "
                    "folder. Three days of prep become three hours."
                ),
                "lessons": [
                    ("Encoding procurement policy as testable rules",
                     "Convert written policies into structured rules Claude can test against transaction data."),
                    ("Approval authority compliance",
                     "Test every transaction against the approval authority matrix and flag violations."),
                    ("Preferred supplier compliance",
                     "Check whether purchases went to preferred suppliers and quantify off-contract spend."),
                    ("Documentation completeness check",
                     "Verify every transaction has the required documents: PO, GRN, invoice, approval record."),
                    ("The compliance ledger",
                     "Build an append-only ledger that logs every compliance test result with timestamp and evidence."),
                    ("Audit package assembly",
                     "Package all evidence (ledger, test results, exceptions, supporting docs) into a structured audit folder."),
                ],
            },
        ],
    },
]


def build_sheet_1_overview(wb):
    ws = wb.active
    ws.title = "Course Overview"
    ws.sheet_properties.tabColor = NAVY

    for col, w in {"A": 10, "B": 10, "C": 52, "D": 50, "E": 55}.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:E1")
    ws["A1"] = "ProcureAI Engineering: The Complete Source-to-Pay Foundation with Claude Code"
    ws["A1"].font = Font("Calibri", 14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    ws.merge_cells("A2:E2")
    ws["A2"] = "5 chapters | 20 courses | 118 lessons | From first prompt to full procurement automation"
    ws["A2"].font = SMALL
    ws["A2"].alignment = Alignment(horizontal="center")

    hdr(ws, 4, ["Chapter", "Course", "Course Name", "What You Will Learn", "How You Will Use It in Procurement"])

    row = 5
    for ch in chapters:
        ch_fill = CH_FILLS[ch["ch"]]
        ch_start = row
        for ci, c in enumerate(ch["courses"]):
            cell(ws, row, 1, f"Ch {ch['ch']}" if ci == 0 else "", CH_FONT, ch_fill, CENTER)
            cell(ws, row, 2, f"{c['num']:02d}", BOLD, None, CENTER)
            cell(ws, row, 3, c["name"], BOLD, None, WRAP)
            cell(ws, row, 4, c["what_you_learn"], NORM, None, WRAP)
            cell(ws, row, 5, c["procurement_use"], NORM, None, WRAP)
            row += 1

        # Chapter summary row
        cell(ws, row, 1, "", font=SMALL, fill=ch_fill)
        ws.merge_cells(start_row=row, start_column=2, end_row=row, end_column=3)
        cell(ws, row, 2, f"Chapter {ch['ch']}: {ch['title']}", CH_FONT, ch_fill, WRAP)
        ws.merge_cells(start_row=row, start_column=4, end_row=row, end_column=5)
        cell(
            ws, row, 4,
            f"What students learn: {ch['what_students_learn']}\n\n"
            f"Procurement application: {ch['procurement_application']}",
            SMALL, ch_fill, WRAP,
        )
        row += 1

        if len(ch["courses"]) > 1:
            ws.merge_cells(start_row=ch_start, start_column=1, end_row=row - 2, end_column=1)

    ws.freeze_panes = "A5"


def build_sheet_2_lessons(wb):
    ws = wb.create_sheet("Lesson Details")
    ws.sheet_properties.tabColor = "2C5F8A"

    for col, w in {"A": 10, "B": 10, "C": 52, "D": 10, "E": 42, "F": 60}.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:F1")
    ws["A1"] = "All 118 Lessons: What You Learn and Why It Matters in Procurement"
    ws["A1"].font = Font("Calibri", 14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    hdr(ws, 3, ["Chapter", "Course", "Course Name", "Lesson", "Lesson Title", "What You Learn and Why It Matters"])

    row = 4
    for ch in chapters:
        ch_fill = CH_FILLS[ch["ch"]]
        for c in ch["courses"]:
            for li, (ltitle, ldesc) in enumerate(c["lessons"]):
                cell(ws, row, 1, f"Ch {ch['ch']}", NORM, ch_fill, CENTER)
                cell(ws, row, 2, f"{c['num']:02d}", NORM, None, CENTER)
                cell(ws, row, 3, c["name"] if li == 0 else "", BOLD if li == 0 else NORM, None, WRAP)
                cell(ws, row, 4, f"L{li + 1}", NORM, None, CENTER)
                cell(ws, row, 5, ltitle, NORM, None, WRAP)
                cell(ws, row, 6, ldesc, NORM, None, WRAP)
                row += 1

    ws.freeze_panes = "A4"


def build_sheet_3_chapters(wb):
    ws = wb.create_sheet("Chapter Summary")
    ws.sheet_properties.tabColor = "4472C4"

    for col, w in {"A": 10, "B": 44, "C": 60, "D": 60, "E": 10, "F": 10}.items():
        ws.column_dimensions[col].width = w

    ws.merge_cells("A1:F1")
    ws["A1"] = "Chapter Summary: Learning Progression from Foundation to Domain Mastery"
    ws["A1"].font = Font("Calibri", 14, bold=True, color=NAVY)
    ws["A1"].alignment = Alignment(horizontal="center")

    hdr(ws, 3, ["Chapter", "Chapter Title", "What Students Learn", "How It Applies to Procurement", "Courses", "Lessons"])

    for ch in chapters:
        r = 3 + ch["ch"]
        ch_fill = CH_FILLS[ch["ch"]]
        total_lessons = sum(len(c["lessons"]) for c in ch["courses"])
        cell(ws, r, 1, f"Ch {ch['ch']}", CH_FONT, ch_fill, CENTER)
        cell(ws, r, 2, ch["title"], BOLD, ch_fill, WRAP)
        cell(ws, r, 3, ch["what_students_learn"], NORM, None, WRAP)
        cell(ws, r, 4, ch["procurement_application"], NORM, None, WRAP)
        cell(ws, r, 5, len(ch["courses"]), BOLD, None, CENTER)
        cell(ws, r, 6, total_lessons, BOLD, None, CENTER)

    tr = 9
    ws.merge_cells(f"A{tr}:B{tr}")
    cell(ws, tr, 1, "TOTAL", BOLD, None, Alignment(horizontal="right", vertical="top"))
    ws.cell(tr, 2).border = THIN
    cell(ws, tr, 3, "From zero to full procurement automation with Claude Code", SMALL, None, WRAP)
    cell(ws, tr, 4, "", NORM)
    cell(ws, tr, 5, 20, BOLD, None, CENTER)
    cell(ws, tr, 6, 118, BOLD, None, CENTER)

    ws.freeze_panes = "A4"


def main():
    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    build_sheet_1_overview(wb)
    build_sheet_2_lessons(wb)
    build_sheet_3_chapters(wb)
    wb.save(str(OUTPUT_PATH))
    print(f"Saved: {OUTPUT_PATH}")
    print(f"Sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
