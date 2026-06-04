"""
Update Course 02 handout to match Course 03 baseline standard.

Changes:
1. Fix bare "Claude" references (add "Claude Code" qualifier).
2. Add "What to learn from this" after each Day-in-the-Life scenario.
3. Add "Why it matters" to Section 3 (folder contents), Section 5 (mental model), Section 10 (pattern).
4. Expand Section 3 with per-folder explanations.
5. Add summary table to Day-in-the-Life section.
"""

import re
import copy
from pathlib import Path
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.enum.text import WD_ALIGN_PARAGRAPH

SCRIPT_DIR = Path(__file__).resolve().parent
PROJECT_ROOT = SCRIPT_DIR.parent
SRC = PROJECT_ROOT / "Handouts" / "Course_02_The_Context_Architect_Handout.docx"
DST = PROJECT_ROOT / "Handouts" / "Course_02_The_Context_Architect_Handout.docx"


def find_para(doc, snippet):
    """Return the first paragraph whose text contains snippet."""
    for p in doc.paragraphs:
        if snippet in p.text:
            return p
    raise ValueError(f"Could not find paragraph containing: {snippet!r}")


def insert_paragraph_after(target_para, text, style_name="Normal", bold_prefix=None):
    """Insert a new paragraph immediately after target_para. Returns the new element."""
    new_p = OxmlElement("w:p")

    # Copy style from a reference or set explicitly
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), style_name)
    pPr.append(pStyle)
    new_p.append(pPr)

    if bold_prefix:
        # Bold run for the prefix
        run_b = OxmlElement("w:r")
        rPr_b = OxmlElement("w:rPr")
        b_elem = OxmlElement("w:b")
        rPr_b.append(b_elem)
        run_b.append(rPr_b)
        t_b = OxmlElement("w:t")
        t_b.set(qn("xml:space"), "preserve")
        t_b.text = bold_prefix
        run_b.append(t_b)
        new_p.append(run_b)

        # Normal run for the rest
        run_n = OxmlElement("w:r")
        t_n = OxmlElement("w:t")
        t_n.set(qn("xml:space"), "preserve")
        t_n.text = text
        run_n.append(t_n)
        new_p.append(run_n)
    else:
        run = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set(qn("xml:space"), "preserve")
        t.text = text
        run.append(t)
        new_p.append(run)

    target_para._element.addnext(new_p)
    return new_p


def insert_heading_after(target_para, text, level=3):
    """Insert a heading paragraph after target_para."""
    new_p = OxmlElement("w:p")
    pPr = OxmlElement("w:pPr")
    pStyle = OxmlElement("w:pStyle")
    pStyle.set(qn("w:val"), f"Heading{level}")
    pPr.append(pStyle)
    new_p.append(pPr)

    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    b_elem = OxmlElement("w:b")
    rPr.append(b_elem)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    new_p.append(run)

    target_para._element.addnext(new_p)
    return new_p


def insert_table_after(doc, target_para, headers, rows):
    """Insert a table after target_para."""
    # Create the table using doc API, then move it
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    tbl.style = "Table Grid"

    # Header row
    for j, h in enumerate(headers):
        cell = tbl.rows[0].cells[j]
        cell.text = h
        for run in cell.paragraphs[0].runs:
            run.bold = True

    # Data rows
    for i, row_data in enumerate(rows):
        for j, val in enumerate(row_data):
            tbl.rows[i + 1].cells[j].text = val

    # Move table element after target paragraph
    target_para._element.addnext(tbl._tbl)
    return tbl


def fix_bare_claude(doc):
    """Replace bare 'Claude reads/finds/writes/...' with 'Claude Code reads/...'."""
    verbs = (
        "reads", "writes", "finds", "opens", "loads", "saves", "checks",
        "reports", "produces", "generates", "confirms", "prints", "runs",
        "follows", "cannot", "will,", "will ", "writes something",
        "sees", "applied",
    )
    count = 0
    for para in doc.paragraphs:
        for run in para.runs:
            if not run.text:
                continue
            original = run.text
            for verb in verbs:
                # Match "Claude <verb>" but NOT "Claude Code <verb>"
                pattern = r"(?<!Claude Code )(?<!Claude Code\n)Claude (" + re.escape(verb) + r")"
                # Only replace if "Claude Code" is not already there
                new_text = run.text
                while True:
                    match = re.search(pattern, new_text)
                    if not match:
                        break
                    # Check it's not already "Claude Code"
                    start = match.start()
                    preceding = new_text[max(0, start - 12):start]
                    if "Claude Code " in preceding:
                        break
                    new_text = new_text[:start] + "Claude Code " + match.group(1) + new_text[match.end():]
                run.text = new_text
            if run.text != original:
                count += 1
    return count


def main():
    doc = Document(str(SRC))

    # ── 1. Fix bare "Claude" references ──
    fixed = fix_bare_claude(doc)
    print(f"Fixed {fixed} bare Claude references.")

    # ── 2. "What to learn from this" after each Day-in-the-Life scenario ──

    # 08:45 scenario ends at "Time: 30 minutes. Without the stack:"
    p = find_para(doc, "Time: 30 minutes. Without the stack:")
    insert_paragraph_after(
        p,
        " The same prompt, run from three different folders, produces three correctly scoped briefs. "
        "The folder you start in determines which CLAUDE.md files load, which determines which data "
        "Claude Code reads. You do not need three different prompts. You need one prompt and three folders.",
        bold_prefix="What to learn from this.",
    )

    # 10:15 scenario ends at "Identified the new rows"
    p = find_para(doc, "Identified the new rows by checking")
    insert_paragraph_after(
        p,
        " Your CLAUDE.md points at files by name, not by content. When the underlying data changes, "
        "the next session picks up the new data automatically. The only maintenance is updating the "
        "row count and the 'last reviewed' date in your CLAUDE.md so the description stays accurate.",
        bold_prefix="What to learn from this.",
    )

    # 11:00 scenario ends at "Returned the violations sorted"
    p = find_para(doc, "Returned the violations sorted by amount descending")
    insert_paragraph_after(
        p,
        " Business rules live in the CLAUDE.md, not in your prompts. When a rule changes, you edit "
        "one file and every future session enforces the new rule. You do not need to remember to "
        "change your prompts or tell Claude Code about the new rule in every conversation.",
        bold_prefix="What to learn from this.",
    )

    # 14:00 scenario ends at "The stack is portable."
    p = find_para(doc, "The stack is portable. Onboarding:")
    insert_paragraph_after(
        p,
        " The CLAUDE.md stack is portable. A new team member does not need training on your prompts "
        "or your data. They get the folder, start Claude Code, and the right context loads "
        "automatically. The five-question test confirms the stack works without running any real analysis.",
        bold_prefix="What to learn from this.",
    )

    # 16:00 scenario ends at "Wrote a one-page profile"
    p = find_para(doc, "Wrote a one-page profile with the figures")
    insert_paragraph_after(
        p,
        " The CLAUDE.md stack supports ad-hoc queries, not just the standard brief. Because Claude "
        "Code already knows the scope, the data files, and the rules for this category, any prompt "
        "you type benefits from that context. A supplier profile, a spend check, a quick risk "
        "flag: all of them use the same stack without extra setup.",
        bold_prefix="What to learn from this.",
    )

    # ── 3. Summary table after "End of Anwar's Tuesday" ──
    p = find_para(doc, "End of Anwar's Tuesday")
    if not p:
        p = find_para(doc, "End of Anwar")
    insert_table_after(
        doc, p,
        headers=["Task", "Time with the stack", "Time without the stack"],
        rows=[
            ["Three category briefs for board pack", "30 minutes", "9 to 12 hours"],
            ["Confirm new suppliers loaded", "2 minutes", "15 minutes"],
            ["Re-run violations after threshold change", "3 minutes", "1 to 2 hours"],
            ["Onboard a new analyst", "30 minutes", "Half a day"],
            ["Supplier meeting prep", "3 minutes", "30 minutes"],
        ],
    )

    # ── 4. Expand Section 3 with per-folder explanations and "Why it matters" ──
    # Insert after "About 10,500 rows of data total..."
    anchor = find_para(doc, "About 10,500 rows of data total")

    # Build the section 3 expansions in REVERSE order (since addnext pushes down)
    section3_items = [
        (
            "scripts/build_course_data.py",
            "A Python script that regenerates all practice data deterministically (it uses a fixed "
            "random seed). If you accidentally delete or corrupt a file during practice, run "
            "`python scripts/build_course_data.py` once and every CSV is restored exactly as it was.",
            "Why this matters. Beginners often accidentally overwrite or delete practice files. One "
            "command restores everything to its original state. You never need to re-download the course.",
        ),
        (
            "solutions/",
            "Reference CLAUDE.md files for the global and each category. These are the answer key. "
            "Look at them only after you have made your own attempt. They are useful for comparing your "
            "sections, your file references, and your rules against a working reference.",
            "Why this matters. Reference answers let you check your own work. If your five-question test "
            "returns different answers from the solution, you know exactly where to fix your CLAUDE.md.",
        ),
        (
            "practice/indirect/",
            "The indirect spend category folder. Contains vendors.csv (80 vendors), invoices.csv "
            "(3,000 invoice rows with 18 approval-rule violations), and a stub CLAUDE.md you will "
            "complete in Lesson 3.",
            "Why this matters. The 18 violations are the key test: when your CLAUDE.md states the $25,000 "
            "approval rule, Claude Code finds exactly 18 invoices that break it. If Claude Code finds a "
            "different number, the rule in your CLAUDE.md is wrong.",
        ),
        (
            "practice/logistics/",
            "The logistics category folder. Contains carriers.csv (20 carriers), shipments.csv "
            "(5,000 shipment rows), and a stub CLAUDE.md you will complete in Lesson 3.",
            "Why this matters. Logistics uses different entity names (carriers, shipments) and different "
            "data files from direct materials (suppliers, orders). When you run the same brief prompt "
            "from this folder, Claude Code produces a logistics brief with carrier names only. That "
            "proves the category CLAUDE.md is working.",
        ),
        (
            "practice/direct-materials/",
            "The direct materials category folder. Contains suppliers.csv (50 suppliers), orders.csv "
            "(2,500 PO rows over the last year), and a stub CLAUDE.md you will complete in Lesson 3.",
            "Why this matters. The category folder isolates one part of your portfolio. When you start "
            "Claude Code here, it reads both the global and the category CLAUDE.md. It sees only the "
            "suppliers, orders, and rules that belong to direct materials. Nothing from logistics or "
            "indirect leaks in.",
        ),
        (
            "practice/CLAUDE.md",
            "The global context file Claude Code reads automatically when you start a session anywhere "
            "inside the practice/ folder. It tells Claude Code who you are (Senior Category Manager), "
            "what the folder rules are, and what writing rules to follow. In Lesson 2 you replace the "
            "starter content with a proper three-section global.",
            "Why this matters. Without this file, Claude Code produces generic output that mixes "
            "categories and misapplies rules. With it, every session starts with the right identity, "
            "the right folder rules, and the right writing standards, before you type a single prompt.",
        ),
    ]

    for heading_text, description, why_text in section3_items:
        # Insert in reverse: why, description, heading
        # Each addnext puts it right after anchor, so reverse order = correct final order
        e3 = insert_paragraph_after(anchor, " " + why_text, bold_prefix="")
        e2 = insert_paragraph_after(anchor, description)
        e1 = insert_heading_after(anchor, heading_text, level=3)

    # ── 5. "Why it matters" for Section 5 (mental model) ──
    p = find_para(doc, "That is the entire mechanism.")
    insert_paragraph_after(
        p,
        " The stack means you write shared context once (the global) and category-specific context "
        "once per folder. You do not repeat yourself. When a writing rule changes, you edit the "
        "global and every category gets the update. When a category-specific rule changes, you edit "
        "one file and the other categories are untouched. One change, one place, every time.",
        bold_prefix="Why this matters.",
    )

    # ── 6. "Why it matters" for Section 10 (the pattern) ──
    p = find_para(doc, "Their personal preferences come from")
    insert_paragraph_after(
        p,
        " A single monolithic CLAUDE.md grows to hundreds of lines and mixes logistics rules with "
        "materials rules. Claude Code reads all of it every session, which means rules from one "
        "category can bleed into another. Small, focused files keep each session clean and make "
        "maintenance easy: change one rule in one place, and only the right category is affected.",
        bold_prefix="Why this matters.",
    )

    # ── Save ──
    doc.save(str(DST))
    print(f"Saved updated handout to {DST}")


if __name__ == "__main__":
    main()
